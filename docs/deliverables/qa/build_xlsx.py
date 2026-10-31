"""The defence question bank as a filterable spreadsheet: python3 build_xlsx.py <bank.json> <out.xlsx>

bank.json comes from `node build.js <out.docx> <bank.json>`. Sheet 1 holds every question with filters on the domain
and the difficulty; sheet 2 the counts per domain and level. Right-to-left, IBM Plex Sans Arabic, design-system colours.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "IBM Plex Sans Arabic"
NAVY, INK, LINE, ZEBRA = "1F3FA0", "182038", "E3E7EF", "F3F5FA"
LEVEL_FILL = {1: "E6F4EA", 2: "E8EDFA", 3: "FFF3D6", 4: "FDE8E8"}

src, out = sys.argv[1], sys.argv[2]
bank = json.load(open(src, encoding="utf-8"))
levels = {int(k): v for k, v in bank["levels"].items()}
rows = bank["rows"]

side = Side(style="thin", color=LINE)
border = Border(left=side, right=side, top=side, bottom=side)
head_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
head_fill = PatternFill("solid", fgColor=NAVY)
wrap = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
centre = Alignment(wrap_text=True, vertical="top", horizontal="center", readingOrder=2)

wb = Workbook()
ws = wb.active
ws.title = "الأسئلة"
ws.sheet_view.rightToLeft = True
heads = ["#", "الاختصاص", "الصعوبة", "رتبة الصعوبة", "السؤال", "الإجابة النموذجية", "العمق الأكاديمي", "المرجع"]
widths = [7, 26, 10, 8, 40, 70, 60, 26]
ws.append(heads)
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font, c.fill, c.alignment, c.border = head_font, head_fill, centre, border
for k, r in enumerate(rows):
    ws.append([r["id"], f'{r["domainN"]}. {r["domain"]}', levels[r["level"]], r["level"], r["q"], r["model"], r["academic"], r["source"]])
    row = ws.max_row
    for c in ws[row]:
        c.font = Font(name=FONT, size=10, color=INK, bold=c.column == 5)
        c.alignment = centre if c.column in (1, 3, 4) else wrap
        c.border = border
        if k % 2:
            c.fill = PatternFill("solid", fgColor=ZEBRA)
    ws.cell(row, 3).fill = PatternFill("solid", fgColor=LEVEL_FILL[r["level"]])
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{ws.max_row}"

# Counts per domain and level, as formulas over sheet 1 so they follow any edit.
s2 = wb.create_sheet("التوزيع")
s2.sheet_view.rightToLeft = True
s2.append(["#", "الاختصاص"] + [levels[l] for l in (1, 2, 3, 4)] + ["المجموع"])
for c in s2[1]:
    c.font, c.fill, c.alignment, c.border = head_font, head_fill, centre, border
last = ws.max_row
for d in bank["domains"]:
    name = f'{d["n"]}. {d["title"]}'
    r = s2.max_row + 1
    s2.append([d["n"], name] + [f'=COUNTIFS(\'الأسئلة\'!$B$2:$B${last},$B{r},\'الأسئلة\'!$D$2:$D${last},{l})' for l in (1, 2, 3, 4)]
              + [f"=SUM(C{r}:F{r})"])
r = s2.max_row + 1
s2.append(["", "المجموع"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{r - 1})" for c in range(3, 8)])
for row in s2.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=FONT, size=11, color=INK, bold=c.row == r)
        c.alignment, c.border = (wrap if c.column == 2 else centre), border
        if c.row == r:
            c.fill = PatternFill("solid", fgColor="E8EDFA")
for i, w in enumerate([6, 44, 11, 11, 11, 11, 11], 1):
    s2.column_dimensions[get_column_letter(i)].width = w

wb.save(out)
print("written", out, len(rows), "rows")
