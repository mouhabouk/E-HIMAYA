"""Parse the three 2026 deployment programmes (Word) into structured rows."""
import docx, glob, json, re, sys
src, out = sys.argv[1], sys.argv[2]


def cells(r):
    o, prev = [], None
    for c in r.cells:
        if c._tc is prev:
            continue
        prev = c._tc
        o.append(" ".join(c.text.split()))
    return o


def num(s):
    m = re.search(r"\d+", s or "")
    return int(m.group()) if m else None


P = {"road": [], "seasonal": [], "beach": []}
# road rescue points
d = docx.Document(glob.glob(src + "/2b567dd1*.docx")[0])
for r in d.tables[0].rows[2:]:
    c = cells(r)
    if not c[-1].strip().isdigit():
        continue
    hours, coord, mat, hum, place, gov, no = c
    lat, lon = [float(x) for x in re.findall(r"-?\d+\.\d+", coord)[:2]]
    P["road"].append(dict(no=int(no), governorate=gov, place=place, staff=hum, equipment=mat, hours=hours, lat=lat, lon=lon,
                          geo_source="إحداثيات البرنامج"))
# seasonal agri/forest centres
d = docx.Document(glob.glob(src + "/5ada505a*.docx")[0])
for r in d.tables[0].rows[2:]:
    c = cells(r)
    if len(c) < 8 or not c[-1].strip().isdigit():
        continue
    note, hours, sup, mat, hum, place, gov, no = c
    m = re.search(r"(\d{2}/\d{2}/\d{4})", note)
    phone = re.findall(r"\d{8}", sup)
    P["seasonal"].append(dict(no=int(no), governorate=gov, place=place, staff_count=num(hum), equipment=mat, hours=hours,
                              opened=m.group(1) if m else None, supervisor=re.sub(r"\d{8}", "", sup).strip(),
                              supervisor_phone=phone[0] if phone else None))
# beach deployment (daily follow-up 04-07-2026)
d = docx.Document(glob.glob(src + "/cc3b36e4*.docx")[0])
gov_prev = None
for t in d.tables:
    for r in t.rows:
        c = cells(r)
        if len(c) < 7 or c[0].startswith(("الوحدة", "المجموع", "لمجموع")) or "المجموع" in c[0]:
            continue
        tail = c[-5:]
        txt = c[:-5]
        sup, deployed, planned, eq, hours = tail
        if not re.search(r"\d", sup + deployed + planned):
            continue
        gov = txt[0] or gov_prev
        gov_prev = gov
        main = txt[1] if len(txt) > 1 else ""
        subs = " + ".join(x for x in txt[2:] if x and x != "//")
        P["beach"].append(dict(governorate=gov, main_point=main, sub_points=subs, supervisors=num(sup), lifeguards_deployed=num(deployed),
                               lifeguards_planned=num(planned), equipment=eq.replace("***", "").strip(), hours=hours))
json.dump(P, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for k, v in P.items():
    print(k, len(v))
from collections import Counter
print(Counter(x["governorate"] for x in P["beach"]))
print(sum(x["lifeguards_deployed"] or 0 for x in P["beach"]), sum(x["lifeguards_planned"] or 0 for x in P["beach"]), sum(x["supervisors"] or 0 for x in P["beach"]))
