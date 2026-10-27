import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
D = json.load(open(src, encoding="utf-8"))
rows, summary = D["rows"], D["summary"]
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); B = Font(name=F); BB = Font(name=F, bold=True)
NAVY = PatternFill("solid", fgColor="1B2A52"); IN = PatternFill("solid", fgColor="FFF2CC")
thin = Border(*(Side(style="thin", color="C9D4EE"),) * 4)
CF = {"مرتفعة": "C6EFCE", "متوسطة": "DDEBF7", "منخفضة": "FCE4D6", "للمراجعة": "F8CBAD"}
LAYER_AR = {s[0]: s[1] for s in summary}

wb = Workbook()
ws = wb.active; ws.title = "الملخص"; ws.sheet_view.rightToLeft = True
ws["A1"] = "استكمال أسماء عناصر OpenStreetMap غير المسمّاة — e-HIMAYA"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = ("لقطة OSM بتاريخ 27-10-2026 · الأسماء المقترحة وصفية (نوع المرفق + أقرب تجمّع سكني أو عنوان) ما لم يُذكر مصدر OSM. "
            "لا يُنشر أي اسم مقترح للمواطن قبل قبوله من مدير الطبقات.")
ws["A2"].font = Font(name=F, italic=True, color="555555")
hdr = ["الطبقة (ملف)", "الطبقة", "كل العناصر", "بدون اسم", "مرتفعة", "متوسطة", "منخفضة", "للمراجعة", "تمت مراجعتها", "% مراجعة"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(4, j, h); c.font = H; c.fill = NAVY; c.alignment = Alignment(horizontal="center", wrap_text=True)
n = len(rows) + 1
r = 5
for s in summary:
    if s[3] == 0: continue
    ws.cell(r, 1, s[0]).font = B; ws.cell(r, 2, s[1]).font = B; ws.cell(r, 3, s[2]).font = Font(name=F, color="0000FF")
    ws.cell(r, 4, f'=COUNTIF(المراجعة!$A$2:$A${n},A{r})')
    for j, k in zip(range(5, 9), ["مرتفعة", "متوسطة", "منخفضة", "للمراجعة"]):
        ws.cell(r, j, f'=COUNTIFS(المراجعة!$A$2:$A${n},$A{r},المراجعة!$K$2:$K${n},"{k}")')
    ws.cell(r, 9, f'=COUNTIFS(المراجعة!$A$2:$A${n},$A{r},المراجعة!$Q$2:$Q${n},"<>")')
    ws.cell(r, 10, f'=IF(D{r}=0,0,I{r}/D{r})').number_format = "0%"
    for j in range(1, 11): ws.cell(r, j).border = thin; ws.cell(r, j).font = ws.cell(r, j).font.copy(name=F)
    r += 1
ws.cell(r, 2, "المجموع").font = BB
for j in range(3, 10):
    L = get_column_letter(j); ws.cell(r, j, f"=SUM({L}5:{L}{r-1})").font = BB
ws.cell(r, 10, f"=IF(D{r}=0,0,I{r}/D{r})").number_format = "0%"
ws.cell(r + 2, 1, "سلّم الثقة").font = BB
leg = [("مرتفعة", "اسم من OSM نفسه: name:en، أو عنصر مسمّى من نفس الطبقة على ≤30 م (تكرار محتمل)، أو مقطع خطي يتصل من الطرفين بمقطع مسمّى"),
       ("متوسطة", "تسمية وصفية: أقرب تجمّع ≤1 كم، أو عنوان OSM، أو عنصر مسمّى قريب (30–100 م)، أو رافد يصبّ في وادٍ مسمّى"),
       ("منخفضة", "أقرب تجمّع بين 1 و3 كم"), ("للمراجعة", "أقرب تجمّع أبعد من 3 كم أو موقع بحري/صحراوي")]
for i, (k, v) in enumerate(leg):
    a = ws.cell(r + 3 + i, 1, k); a.fill = PatternFill("solid", fgColor=CF[k]); a.font = BB
    ws.cell(r + 3 + i, 2, v).font = B
ws.cell(r + 8, 1, "طريقة الاستعمال").font = BB
ws.cell(r + 9, 1, "في ورقة «المراجعة»: الخانات الصفراء (Q قرار المدير، R الاسم المعتمد، S ملاحظة) هي وحدها للتعبئة. القرار «قبول» يعتمد الاسم المقترح، و«تعديل» يعتمد ما في R، و«رفض» يُبقي العنصر دون اسم ظاهر.").font = B
ws.cell(r + 10, 1, "الطبقات الحساسة (ثكنات، النفط والغاز) معلّمة في العمود O ولا تُعرض إلا للوزير حسب قاعدة الظهور المتفق عليها.").font = B
for col, w in zip("ABCDEFGHIJ", [18, 26, 11, 10, 10, 10, 10, 10, 12, 10]): ws.column_dimensions[col].width = w

wr = wb.create_sheet("المراجعة"); wr.sheet_view.rightToLeft = True
cols = ["الطبقة (ملف)", "الطبقة", "osm_id", "النوع", "الولاية", "lat", "lon", "الصنف", "الاسم المقترح", "مصدر الاقتراح", "درجة الثقة",
        "أقرب تجمّع", "المسافة كم", "تكرار محتمل", "حساسية", "رابط OSM", "قرار المدير", "الاسم المعتمد", "ملاحظة"]
for j, h in enumerate(cols, 1):
    c = wr.cell(1, j, h); c.font = H; c.fill = NAVY; c.alignment = Alignment(horizontal="center", wrap_text=True)
for i, row in enumerate(rows, 2):
    vals = [row[0], LAYER_AR.get(row[0], row[1])] + row[2:]
    for j, v in enumerate(vals, 1):
        c = wr.cell(i, j, v); c.font = B
    for j in (17, 18, 19): wr.cell(i, j).fill = IN
    wr.cell(i, 16).hyperlink = row[15]
dv = DataValidation(type="list", formula1='"قبول,تعديل,رفض"', allow_blank=True)
wr.add_data_validation(dv); dv.add(f"Q2:Q{n}")
for k, col in CF.items():
    wr.conditional_formatting.add(f"K2:K{n}", CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=col)))
wr.freeze_panes = "C2"; wr.auto_filter.ref = f"A1:S{n}"
for j, w in enumerate([14, 22, 13, 6, 12, 9, 9, 18, 40, 38, 11, 20, 9, 9, 8, 38, 11, 30, 24], 1):
    wr.column_dimensions[get_column_letter(j)].width = w
wb.save(out)
print("rows", len(rows))
