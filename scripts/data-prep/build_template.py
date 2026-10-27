"""Build the deployment-programme import template (xlsx + csv + json + schema), pre-filled with the 2026 programmes."""
import csv, json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

D = sys.argv[1]
P = json.load(open(f"{D}/programs_2026.json", encoding="utf-8"))
G = {}
for r in open(f"{D}/geo_results.txt", encoding="utf-8").read().strip().split(";"):
    k, lat, lon, kind, sc, by = r.split(",")
    G[k] = (float(lat) if lat else None, float(lon) if lon else None, kind, float(sc), by)
KIND = {"b": "شاطئ مسمّى", "v": "قرية", "t": "بلدة", "c": "مدينة", "n": "حي", "s": "حي", "h": "تجمّع", "l": "موقع"}


def conf(key):
    lat, lon, kind, sc, by = G[key]
    if lat is None:
        return None, None, "غير محدد: يحدده المدير على الخريطة"
    if kind == "b":
        return lat, lon, "مرتفعة: شاطئ مسمّى في OSM"
    if sc >= 3 and by == "s":
        return lat, lon, "متوسطة: النقطة الفرعية (" + KIND.get(kind, kind) + ")"
    if sc >= 3:
        return lat, lon, "تقريبية: مركز " + KIND.get(kind, kind) + " (البلدة لا النقطة)"
    return lat, lon, "للمراجعة: مطابقة جزئية للاسم"


PROGRAMS = [
    dict(program_id="ROAD-2026", program_type="نقاط النجدة والإسعاف بالطرقات", title="مراكز النجدة والإسعاف بالطرقات لتأمين الموسم الصيفي والسياحي",
         season="صائفة 2026", valid_from="2026-06-15", valid_to="2026-09-15", version=1, status="معتمد", approved_by="", approval_date="",
         source_document="جدول مراكز النجدة والإسعاف بالطرقات لصائفة 2026 (Word)"),
    dict(program_id="SEAS-2026", program_type="المراكز الموسمية", title="تركيز المراكز الموسمية لحماية الثروة الفلاحية والغابية (حماية مدنية وغابات)",
         season="صائفة 2026", valid_from="2026-06-01", valid_to="2026-09-30", version=1, status="معتمد", approved_by="", approval_date="",
         source_document="جدول تركيز المراكز الموسمية لصائفة 2026 (Word)"),
    dict(program_id="BEACH-2026", program_type="التركز بالشواطئ", title="التركيز التدريجي لوحدات الحماية المدنية بالشواطئ",
         season="صائفة 2026", valid_from="2026-06-15", valid_to="2026-09-15", version=1, status="معتمد", approved_by="", approval_date="",
         source_document="جدول متابعة التركيز التدريجي بالشواطئ ليوم 04-07-2026 (Word)"),
]
POINT_COLS = [
    ("program_id", "معرّف البرنامج", "إلزامي، من ورقة البرامج"),
    ("point_id", "معرّف النقطة", "إلزامي وفريد داخل البرنامج"),
    ("governorate", "الولاية", "إلزامي، من القائمة"),
    ("place_name", "المكان / النقطة الرئيسية", "إلزامي"),
    ("sub_points", "النقاط الفرعية", "اختياري، تفصل بعلامة +"),
    ("latitude", "خط العرض", "عشري WGS84، مثل 36.85"),
    ("longitude", "خط الطول", "عشري WGS84، مثل 10.32"),
    ("geo_confidence", "دقة التموقع", "تُملأ آلياً عند الاستيراد"),
    ("staff_count", "عدد الأعوان", "عدد صحيح"),
    ("staff_composition", "تركيبة الفريق", "نص حر"),
    ("supervisors", "عدد المشرفين", "عدد صحيح"),
    ("lifeguards_planned", "السباحون المنقذون (المبرمج)", "عدد صحيح، للشواطئ"),
    ("lifeguards_deployed", "السباحون المنقذون (المركّز)", "عدد صحيح، يُحيَّن يومياً"),
    ("equipment", "الإمكانيات المادية", "نص حر"),
    ("hours", "توقيت العمل", "مثل 24/24 أو 10:00-18:00"),
    ("opening_date", "تاريخ الفتح", "YYYY-MM-DD"),
    ("closing_date", "تاريخ الغلق", "YYYY-MM-DD"),
    ("point_status", "حالة النقطة", "مبرمجة / مفتوحة / مغلقة"),
    ("supervisor_name", "المشرف (مقيّد)", "لا يظهر إلا لأعوان القاعة"),
    ("supervisor_phone", "هاتف المشرف (مقيّد)", "لا يظهر إلا لأعوان القاعة"),
    ("notes", "ملاحظات", "اختياري"),
]
PROG_COLS = [("program_id", "معرّف البرنامج"), ("program_type", "نوع البرنامج"), ("title", "العنوان"), ("season", "الموسم"),
             ("valid_from", "ساري من"), ("valid_to", "ساري إلى"), ("version", "النسخة"), ("status", "الحالة"),
             ("approved_by", "اعتمده"), ("approval_date", "تاريخ الاعتماد"), ("source_document", "الوثيقة المصدر")]

points = []
for s in P["road"]:
    points.append(dict(program_id="ROAD-2026", point_id=f"ROAD-2026-{s['no']:02d}", governorate=s["governorate"], place_name=s["place"],
                       latitude=s["lat"], longitude=s["lon"], geo_confidence="مؤكدة: إحداثيات البرنامج", staff_composition=s["staff"],
                       equipment=s["equipment"], hours=s["hours"], point_status="مفتوحة"))
for i, s in enumerate(P["seasonal"]):
    lat, lon, c = conf(f"S{i}")
    d = s["opened"]
    iso = f"{d[6:]}-{d[3:5]}-{d[:2]}" if d else None
    points.append(dict(program_id="SEAS-2026", point_id=f"SEAS-2026-{s['no']:02d}", governorate=s["governorate"], place_name=s["place"],
                       latitude=lat, longitude=lon, geo_confidence=c, staff_count=s["staff_count"], equipment=s["equipment"], hours=s["hours"],
                       opening_date=iso, point_status="مفتوحة", supervisor_name=s["supervisor"], supervisor_phone=s["supervisor_phone"]))
for i, b in enumerate(P["beach"]):
    lat, lon, c = conf(f"B{i}")
    points.append(dict(program_id="BEACH-2026", point_id=f"BEACH-2026-{i+1:03d}", governorate=b["governorate"], place_name=b["main_point"],
                       sub_points=b["sub_points"], latitude=lat, longitude=lon, geo_confidence=c, supervisors=b["supervisors"],
                       lifeguards_planned=b["lifeguards_planned"], lifeguards_deployed=b["lifeguards_deployed"], equipment=b["equipment"],
                       hours=b["hours"], point_status="مفتوحة", notes="الأعداد من جدول المتابعة ليوم 04-07-2026"))

GOVS = ["أريانة", "باجة", "بن عروس", "بنزرت", "تطاوين", "توزر", "تونس", "جندوبة", "زغوان", "سليانة", "سوسة", "سيدي بوزيد", "صفاقس",
        "قابس", "قبلي", "القصرين", "قفصة", "القيروان", "الكاف", "مدنين", "المنستير", "المهدية", "منوبة", "نابل"]
TYPES = ["نقاط النجدة والإسعاف بالطرقات", "المراكز الموسمية", "التركز بالشواطئ", "برنامج آخر"]
PSTAT = ["مسودة", "معتمد", "مُحيَّن", "منتهٍ", "ملغى"]
QSTAT = ["مبرمجة", "مفتوحة", "مغلقة"]

F = "Arial"
hdr_fill = PatternFill("solid", fgColor="1B2A52")
inp_fill = PatternFill("solid", fgColor="FFF2CC")
ro_fill = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

# instructions
ws = wb.active
ws.title = "تعليمات"
ws.sheet_view.rightToLeft = True
lines = [
    ("نموذج برامج التركيز الموسمية لمنظومة ء-حماية", True),
    ("يُستعمل لإدخال البرامج السنوية أو الفصلية (نقاط النجدة بالطرقات، المراكز الموسمية، التركز بالشواطئ، أو أي برنامج آخر) ورفعها إلى المنظومة لتظهر طبقاتٍ على الخريطة.", False),
    ("", False),
    ("1. ورقة «البرامج»: سطر واحد لكل برنامج. المعرّف فريد (مثل BEACH-2027). تاريخا السريان يحددان متى تظهر الطبقة على الخريطة.", False),
    ("2. ورقة «النقاط»: سطر لكل نقطة، مربوط بالبرنامج عبر «معرّف البرنامج». الخانات الصفراء للإدخال، والرمادية تُملأ آلياً.", False),
    ("3. الإحداثيات اختيارية عند الإدخال: إن غابت يقترحها النظام من الاسم ويؤكدها المدير بالنقر على الخريطة.", False),
    ("4. التحيين: عدّل الأسطر وارفع الملف من جديد بنفس معرّف البرنامج ونسخة أعلى؛ النظام يعرض الفروق قبل الاعتماد ويحفظ النسخة السابقة.", False),
    ("5. خانتا المشرف وهاتفه مقيّدتان: لا تظهران إلا لأعوان القاعة المعنية (معطيات شخصية، القانون الأساسي عدد 63 لسنة 2004).", False),
    ("6. الصيغ المقبولة للرفع: هذا الملف (xlsx)، أو CSV بنفس أسماء الأعمدة الإنجليزية، أو JSON حسب الملف program_schema.json.", False),
    ("", False),
    ("محتوى الملف الآن: برامج صائفة 2026 الثلاثة، محوّلة من وثائق Word كمثال حقيقي. دقة التموقع مذكورة لكل نقطة في عمود «دقة التموقع».", False),
    ("ملاحظة على المصدر: المجموع العام في جدول الشواطئ (2154 / 1439) لا يطابق مجموع المجاميع الجهوية؛ لذلك تُحسب المجاميع هنا بالمعادلات من الأسطر.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=b, size=14 if b else 11)
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")
ws.column_dimensions["A"].width = 120

# lists
wl = wb.create_sheet("القوائم")
wl.sheet_view.rightToLeft = True
for j, (name, vals) in enumerate([("الولايات", GOVS), ("أنواع البرامج", TYPES), ("حالة البرنامج", PSTAT), ("حالة النقطة", QSTAT)], 1):
    wl.cell(row=1, column=j, value=name).font = Font(name=F, bold=True)
    for i, v in enumerate(vals, 2):
        wl.cell(row=i, column=j, value=v).font = Font(name=F)
    wl.column_dimensions[chr(64 + j)].width = 32

# programmes
wp = wb.create_sheet("البرامج", 1)
wp.sheet_view.rightToLeft = True
extra = [("عدد النقاط", None), ("نقاط بلا إحداثيات", None), ("الأعوان", None), ("المنقذون المبرمجون", None), ("المنقذون المركّزون", None)]
heads = [h for _, h in PROG_COLS] + [h for h, _ in extra]
for j, h in enumerate(heads, 1):
    c = wp.cell(row=1, column=j, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = B
for j, (k, _) in enumerate(PROG_COLS, 1):
    wp.cell(row=2, column=j).comment = None
NP = len(points) + 1
for i, p in enumerate(PROGRAMS, 2):
    for j, (k, _) in enumerate(PROG_COLS, 1):
        c = wp.cell(row=i, column=j, value=p[k] if p[k] != "" else None)
        c.font = Font(name=F, color="0000FF")
        c.fill = inp_fill
        c.border = B
    n = len(PROG_COLS)
    f = [f"=COUNTIF('النقاط'!$A$2:$A${NP},A{i})",
         f"=COUNTIFS('النقاط'!$A$2:$A${NP},A{i},'النقاط'!$F$2:$F${NP},\"\")",
         f"=SUMIFS('النقاط'!$I$2:$I${NP},'النقاط'!$A$2:$A${NP},A{i})",
         f"=SUMIFS('النقاط'!$L$2:$L${NP},'النقاط'!$A$2:$A${NP},A{i})",
         f"=SUMIFS('النقاط'!$M$2:$M${NP},'النقاط'!$A$2:$A${NP},A{i})"]
    for k2, fx in enumerate(f, 1):
        c = wp.cell(row=i, column=n + k2, value=fx)
        c.font = Font(name=F)
        c.fill = ro_fill
        c.border = B
widths = [14, 30, 52, 14, 13, 13, 9, 11, 16, 14, 50, 11, 14, 11, 16, 16]
for j, w in enumerate(widths, 1):
    wp.column_dimensions[wp.cell(row=1, column=j).column_letter].width = w
wp.freeze_panes = "B2"
wp.cell(row=1, column=1).comment = Comment("الخانات الصفراء للإدخال؛ الرمادية معادلات تُحسب من ورقة النقاط.", "ء-حماية")
dv = DataValidation(type="list", formula1=f"='القوائم'!$B$2:$B${len(TYPES)+1}", allow_blank=False)
wp.add_data_validation(dv); dv.add("B2:B200")
dv = DataValidation(type="list", formula1=f"='القوائم'!$C$2:$C${len(PSTAT)+1}", allow_blank=False)
wp.add_data_validation(dv); dv.add("H2:H200")

# points
wq = wb.create_sheet("النقاط", 2)
wq.sheet_view.rightToLeft = True
for j, (k, h, hint) in enumerate(POINT_COLS, 1):
    c = wq.cell(row=1, column=j, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="8A1C1C") if "مقيّد" in h else hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = B
    c.comment = Comment(f"{k}: {hint}", "ء-حماية")
for i, p in enumerate(points, 2):
    for j, (k, _, _) in enumerate(POINT_COLS, 1):
        v = p.get(k)
        c = wq.cell(row=i, column=j, value=v)
        c.font = Font(name=F, color="000000" if k == "geo_confidence" else "0000FF")
        c.fill = ro_fill if k == "geo_confidence" else inp_fill
        c.border = B
        if k in ("latitude", "longitude") and v is not None:
            c.number_format = "0.00000"
pw = [12, 16, 10, 34, 40, 10, 10, 34, 9, 26, 9, 11, 11, 30, 22, 12, 12, 11, 22, 14, 30]
for j, w in enumerate(pw, 1):
    wq.column_dimensions[wq.cell(row=1, column=j).column_letter].width = w
wq.row_dimensions[1].height = 36
wq.freeze_panes = "E2"
wq.auto_filter.ref = f"A1:{wq.cell(row=1, column=len(POINT_COLS)).column_letter}{NP}"
dv = DataValidation(type="list", formula1=f"='القوائم'!$A$2:$A${len(GOVS)+1}", allow_blank=False)
wq.add_data_validation(dv); dv.add(f"C2:C2000")
dv = DataValidation(type="list", formula1=f"='القوائم'!$D$2:$D${len(QSTAT)+1}")
wq.add_data_validation(dv); dv.add(f"R2:R2000")
dv = DataValidation(type="decimal", operator="between", formula1="30", formula2="37.6", allow_blank=True, error="خط العرض خارج حدود تونس")
wq.add_data_validation(dv); dv.add("F2:F2000")
dv = DataValidation(type="decimal", operator="between", formula1="7.4", formula2="11.7", allow_blank=True, error="خط الطول خارج حدود تونس")
wq.add_data_validation(dv); dv.add("G2:G2000")
for s in wb.worksheets:
    s.sheet_properties.tabColor = {"تعليمات": "1B2A52", "البرامج": "E08600", "النقاط": "2B8A3E", "القوائم": "999999"}[s.title]
wb.save(f"{D}/out/e-himaya_programmes_template_2026.xlsx")

# csv + json
with open(f"{D}/out/programme_points_2026.csv", "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=[k for k, _, _ in POINT_COLS])
    w.writeheader()
    for p in points:
        w.writerow({k: p.get(k) for k, _, _ in POINT_COLS})
json.dump({"programmes": [dict(p, points=[{k: v for k, v in q.items() if v not in (None, "")} for q in points if q["program_id"] == p["program_id"]])
                          for p in PROGRAMS]}, open(f"{D}/out/programmes_2026.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
schema = {
    "$schema": "https://json-schema.org/draft/2020-12/schema", "title": "e-HIMAYA deployment programmes", "type": "object",
    "required": ["programmes"],
    "properties": {"programmes": {"type": "array", "items": {
        "type": "object", "required": ["program_id", "program_type", "title", "valid_from", "valid_to", "version", "status", "points"],
        "properties": {
            "program_id": {"type": "string", "pattern": "^[A-Z]+-[0-9]{4}(-[A-Z0-9]+)?$"},
            "program_type": {"enum": TYPES}, "title": {"type": "string"}, "season": {"type": "string"},
            "valid_from": {"type": "string", "format": "date"}, "valid_to": {"type": "string", "format": "date"},
            "version": {"type": "integer", "minimum": 1}, "status": {"enum": PSTAT},
            "approved_by": {"type": "string"}, "approval_date": {"type": "string", "format": "date"}, "source_document": {"type": "string"},
            "points": {"type": "array", "items": {
                "type": "object", "required": ["point_id", "governorate", "place_name"],
                "properties": {
                    "point_id": {"type": "string"}, "governorate": {"enum": GOVS}, "place_name": {"type": "string"},
                    "sub_points": {"type": "string"},
                    "latitude": {"type": "number", "minimum": 30.0, "maximum": 37.6},
                    "longitude": {"type": "number", "minimum": 7.4, "maximum": 11.7},
                    "staff_count": {"type": "integer", "minimum": 0}, "staff_composition": {"type": "string"},
                    "supervisors": {"type": "integer", "minimum": 0},
                    "lifeguards_planned": {"type": "integer", "minimum": 0}, "lifeguards_deployed": {"type": "integer", "minimum": 0},
                    "equipment": {"type": "string"}, "hours": {"type": "string"},
                    "opening_date": {"type": "string", "format": "date"}, "closing_date": {"type": "string", "format": "date"},
                    "point_status": {"enum": QSTAT},
                    "supervisor_name": {"type": "string", "description": "restricted: hall staff only"},
                    "supervisor_phone": {"type": "string", "description": "restricted: hall staff only"},
                    "notes": {"type": "string"}}}}}}}}}
json.dump(schema, open(f"{D}/out/program_schema.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)

from collections import Counter
print(len(points), Counter(p["program_id"] for p in points))
print(Counter(p["geo_confidence"].split(":")[0] for p in points))
