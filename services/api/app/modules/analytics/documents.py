"""Documents of the reports repository (FR-C-21): content, PDF (WeasyPrint, Arabic RTL or French) and XLSX (openpyxl).

A document is a cover (title, kind, period, gouvernorats, filters used, the k-anonymity rule) followed by sections of
tables. Count cells under ``ANALYTICS_MIN_CELL`` are printed "<N" in every export, and the durations computed on fewer
reports than that are left out, so an exported document never singles out a person.
"""

import html
import io
import uuid
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from typing import Any

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.modules.analytics import domain, service
from app.modules.analytics.models import Report, Scenario

Cell = str | int | float | None


@dataclass
class Column:
    label: dict[str, str]
    # Count column: the k-anonymity floor applies to its cells.
    count: bool = False


@dataclass
class Section:
    title: dict[str, str]
    columns: list[Column]
    rows: list[list[Any]] = field(default_factory=list)
    note: dict[str, str] | None = None


@dataclass
class Content:
    title: str
    kind: str
    lang: str
    period: tuple[str, str]
    gouvernorats: list[str]
    filters: list[tuple[dict[str, str], str]]
    min_cell: int
    generated_at: datetime
    sections: list[Section]


class ContentError(Exception):
    """The document cannot be built (its incident or scenario is gone): the report fails with this code."""


def _l(ar: str, fr: str) -> dict[str, str]:
    return {"ar": ar, "fr": fr}


def _minutes(seconds: float | None, n: int, min_cell: int) -> Cell:
    """A duration in minutes, left out when it was computed on fewer reports than the floor."""
    if seconds is None or n < min_cell:
        return None
    return round(seconds / 60.0)


def _counts(title: dict[str, str], key_label: dict[str, str], rows: list[tuple[Any, int]]) -> Section:
    return Section(title, [Column(key_label), Column(_l("العدد", "Nombre"), count=True)], [list(r) for r in rows])


async def _period_sections(
    session: AsyncSession, report: Report, names: dict[int, dict[str, str]], by_gouvernorat: bool
) -> list[Section]:
    min_cell = get_settings().analytics_min_cell
    category = report.filters.get("category")
    data = await service.overview(session, report.gouvernorats, category, report.period_from, report.period_to)
    lang = report.lang
    t = data.totals
    rt = data.response_time
    sections = [
        Section(
            _l("المؤشرات الرئيسية", "Indicateurs clés"),
            [Column(_l("المؤشر", "Indicateur")), Column(_l("القيمة", "Valeur"), count=True)],
            [
                [_l("البلاغات المستلمة", "Signalements reçus")[lang], t.received],
                [_l("البلاغات المؤكدة", "Signalements vérifiés")[lang], t.verified],
                [_l("قيد المعالجة", "En cours")[lang], t.in_progress],
                [_l("مغلقة", "Clôturés")[lang], t.closed],
                [_l("مرفوضة", "Rejetés")[lang], t.rejected],
            ],
        ),
        Section(
            _l("آجال الاستجابة (بالدقائق)", "Délais de réponse (minutes)"),
            [
                Column(_l("المرحلة", "Étape")),
                Column(_l("المتوسط", "Moyenne")),
                Column(_l("المئين 90", "90e percentile")),
                Column(_l("عدد البلاغات", "Signalements"), count=True),
            ],
            [
                [label[lang], _minutes(s.mean_s, s.n, min_cell), _minutes(s.p90_s, s.n, min_cell), s.n]
                for label, s in (
                    (_l("الاستلام ← أول تكليف", "Réception → première affectation"), rt.to_assignment),
                    (_l("الاستلام ← الوصول إلى المكان", "Réception → arrivée sur les lieux"), rt.to_on_scene),
                    (_l("الاستلام ← الإغلاق", "Réception → clôture"), rt.to_closed),
                )
            ],
            note=_l(
                f"لا تُعرض المدد المحسوبة على أقل من {min_cell} بلاغات.",
                f"Les durées calculées sur moins de {min_cell} signalements ne sont pas affichées.",
            ),
        ),
        _counts(
            _l("حسب الصنف", "Par catégorie"),
            _l("الصنف", "Catégorie"),
            [(domain.category_label(str(g.key))[lang], g.count) for g in data.by_category],
        ),
        _counts(
            _l("حسب الحالة", "Par statut"),
            _l("الحالة", "Statut"),
            [(domain.label(domain.STATUS_LABELS, str(g.key))[lang], g.count) for g in data.by_status],
        ),
        _counts(
            _l("حسب المصدر", "Par source"),
            _l("المصدر", "Source"),
            [(domain.label(domain.SOURCE_LABELS, str(g.key))[lang], g.count) for g in data.by_source],
        ),
        _counts(
            _l("حسب الولاية", "Par gouvernorat"),
            _l("الولاية", "Gouvernorat"),
            [
                (names.get(int(g.key), {}).get(lang, "—") if g.key != "unknown" else "—", g.count)
                for g in data.by_gouvernorat
            ],
        ),
        _counts(
            _l("حسب ساعة اليوم", "Par heure"),
            _l("الساعة", "Heure"),
            [(f"{int(g.key):02d}:00", g.count) for g in data.by_hour],
        ),
        _counts(
            _l("حسب يوم الأسبوع", "Par jour de la semaine"),
            _l("اليوم", "Jour"),
            [(domain.WEEKDAY_LABELS[int(g.key)][lang], g.count) for g in data.by_weekday],
        ),
    ]
    if by_gouvernorat:
        sections.append(await _gouvernorat_matrix(session, report, names))
    return sections


async def _gouvernorat_matrix(session: AsyncSession, report: Report, names: dict[int, dict[str, str]]) -> Section:
    """Reports per gouvernorat and category (gouvernorat summary)."""
    params: dict[str, Any] = {
        "t0": domain.day_start(report.period_from),
        "t1": domain.day_start(report.period_to + timedelta(days=1)),
        "codes": report.gouvernorats,
    }
    scope = "AND gouvernorat_code = ANY(CAST(:codes AS smallint[]))" if report.gouvernorats else ""
    if category := report.filters.get("category"):
        scope += " AND category = :category"
        params["category"] = category
    rows = (
        await session.execute(
            text(
                "SELECT gouvernorat_code, category, count(*) FROM reports.report "  # noqa: S608 (constant SQL fragments, values are bound)
                f"WHERE received_at >= :t0 AND received_at < :t1 {scope} GROUP BY 1, 2"
            ),
            params,
        )
    ).all()
    categories = sorted({c for _, c, _ in rows})
    matrix: dict[int | None, dict[str, int]] = {}
    for code, cat, n in rows:
        matrix.setdefault(code, {})[cat] = int(n)
    lang = report.lang
    return Section(
        _l("البلاغات حسب الولاية والصنف", "Signalements par gouvernorat et catégorie"),
        [Column(_l("الولاية", "Gouvernorat"))]
        + [Column(domain.category_label(c), count=True) for c in categories]
        + [Column(_l("المجموع", "Total"), count=True)],
        [
            [names.get(code, {}).get(lang, "—") if code is not None else "—"]
            + [cats.get(c, 0) for c in categories]
            + [sum(cats.values())]
            for code, cats in sorted(matrix.items(), key=lambda kv: -sum(kv[1].values()))
        ],
    )


async def _incident_sections(
    session: AsyncSession, report: Report, names: dict[int, dict[str, str]], after_action: bool
) -> list[Section]:
    lang = report.lang
    incident_id = report.filters.get("incident_id")
    row = (
        await session.execute(
            text(
                "SELECT category, gouvernorat_code, severity, CAST(op_level AS text), status, opened_at, closed_at, "
                "classification, evaluation FROM ops.incident WHERE id = :id"
            ),
            {"id": incident_id},
        )
    ).one_or_none()
    if row is None:
        raise ContentError("incident_not_found")
    category, code, severity, level, state, opened, closed, classification, evaluation = row
    fmt = "%Y-%m-%d %H:%M"
    card = Section(
        _l("بطاقة التدخل", "Fiche d'intervention"),
        [Column(_l("البند", "Élément")), Column(_l("القيمة", "Valeur"))],
        [
            [_l("الصنف", "Catégorie")[lang], domain.category_label(category)[lang]],
            [_l("الولاية", "Gouvernorat")[lang], names.get(code, {}).get(lang, "—") if code else "—"],
            [_l("الخطورة", "Gravité")[lang], severity],
            [_l("المستوى العملياتي", "Niveau opérationnel")[lang], level],
            [_l("التصنيف", "Classification")[lang], classification or "—"],
            [_l("الحالة", "Statut")[lang], state],
            [_l("الفتح", "Ouverture")[lang], opened.astimezone(domain.TUNIS_TZ).strftime(fmt)],
            [_l("الإغلاق", "Clôture")[lang], closed.astimezone(domain.TUNIS_TZ).strftime(fmt) if closed else "—"],
        ],
    )
    params = {"id": incident_id}
    reports_rows = (
        await session.execute(
            text(
                "SELECT category, status, source, count(*) FROM reports.report WHERE incident_id = :id GROUP BY 1, 2, 3"
            ),
            params,
        )
    ).all()
    by: dict[str, dict[str, int]] = {"category": {}, "status": {}, "source": {}}
    for cat, st, src, n in reports_rows:
        for key, value in (("category", cat), ("status", st), ("source", src)):
            by[key][value] = by[key].get(value, 0) + int(n)
    reports_section = Section(
        _l("البلاغات المرتبطة", "Signalements rattachés"),
        [Column(_l("البعد", "Dimension")), Column(_l("القيمة", "Valeur")), Column(_l("العدد", "Nombre"), count=True)],
        [[_l("الصنف", "Catégorie")[lang], domain.category_label(k)[lang], n] for k, n in sorted(by["category"].items())]
        + [
            [_l("الحالة", "Statut")[lang], domain.label(domain.STATUS_LABELS, k)[lang], n]
            for k, n in sorted(by["status"].items())
        ]
        + [
            [_l("المصدر", "Source")[lang], domain.label(domain.SOURCE_LABELS, k)[lang], n]
            for k, n in sorted(by["source"].items())
        ],
    )
    resources = (
        await session.execute(
            text(
                "SELECT r.kind, count(*), "
                "avg(extract(epoch FROM a.on_scene_at - a.created_at)) FILTER (WHERE a.on_scene_at IS NOT NULL), "
                "avg(extract(epoch FROM a.released_at - a.created_at)) FILTER (WHERE a.released_at IS NOT NULL) "
                "FROM ops.assignment a JOIN ops.resource r ON r.id = a.resource_id "
                "WHERE a.incident_id = :id AND a.kind = 'resource' AND a.status <> 'cancelled' GROUP BY 1 ORDER BY 1"
            ),
            params,
        )
    ).all()
    volunteers = await session.scalar(
        text(
            "SELECT count(*) FROM ops.assignment "
            "WHERE incident_id = :id AND kind = 'volunteer' AND status <> 'cancelled'"
        ),
        params,
    )
    resources_section = Section(
        _l("الوسائل المسخرة", "Moyens engagés"),
        [
            Column(_l("النوع", "Type")),
            Column(_l("العدد", "Nombre"), count=True),
            Column(_l("متوسط الوصول (دقيقة)", "Arrivée moyenne (min)")),
            Column(_l("متوسط مدة التسخير (دقيقة)", "Engagement moyen (min)")),
        ],
        [
            [domain.label(domain.RESOURCE_LABELS, kind)[lang], int(n), _rnd(arrive), _rnd(engaged)]
            for kind, n, arrive, engaged in resources
        ]
        + [[_l("متطوعون", "Bénévoles")[lang], int(volunteers or 0), None, None]],
    )
    sections = [card, reports_section, resources_section]
    if after_action:
        timeline = (
            await session.execute(
                text(
                    "SELECT kind, count(*), min(at), max(at) FROM ops.timeline_event WHERE incident_id = :id "
                    "GROUP BY 1 ORDER BY min(at)"
                ),
                params,
            )
        ).all()
        sections.append(
            Section(
                _l("الخط الزمني (حسب نوع الحدث)", "Chronologie (par type d'événement)"),
                [
                    Column(_l("النوع", "Type")),
                    Column(_l("العدد", "Nombre"), count=True),
                    Column(_l("الأول", "Premier")),
                    Column(_l("الأخير", "Dernier")),
                ],
                [
                    [
                        kind,
                        int(n),
                        first.astimezone(domain.TUNIS_TZ).strftime(fmt),
                        last.astimezone(domain.TUNIS_TZ).strftime(fmt),
                    ]
                    for kind, n, first, last in timeline
                ],
            )
        )
        ev = evaluation or {}
        casualties = ev.get("casualties") or {}
        sections.append(
            Section(
                _l("التقييم عند الإغلاق", "Évaluation à la clôture"),
                [Column(_l("البند", "Élément")), Column(_l("القيمة", "Valeur"), count=True)],
                [[_l("النتيجة", "Issue")[lang], ev.get("outcome") or "—"]]
                + [[str(k), int(v)] for k, v in sorted(casualties.items()) if isinstance(v, int)],
                note=_l(
                    "النصوص الحرة للتقييم (الملخص، الدروس) تبقى في ملف التدخل ولا تُنسخ هنا.",
                    "Les textes libres de l'évaluation (synthèse, leçons) restent dans le dossier "
                    "et ne sont pas repris ici.",
                ),
            )
        )
    return sections


def _rnd(seconds: Any) -> int | None:
    return round(float(seconds) / 60.0) if seconds is not None else None


async def _scenario_sections(session: AsyncSession, report: Report) -> list[Section]:
    lang = report.lang
    scenario = await session.get(Scenario, uuid.UUID(str(report.filters.get("scenario_id"))))
    if scenario is None or scenario.result is None:
        raise ContentError("scenario_not_found")
    result, p = scenario.result, scenario.params
    summary = result["summary"]
    prov = result["provenance"]
    band_cols = [
        Column(_l("المؤشر", "Indicateur")),
        Column(_l("م10", "P10")),
        Column(_l("م50", "P50")),
        Column(_l("م90", "P90")),
    ]
    sections = [
        Section(
            _l("المعطيات", "Paramètres"),
            [Column(_l("البند", "Élément")), Column(_l("القيمة", "Valeur"))],
            [
                [_l("الاسم", "Nom")[lang], scenario.name],
                [_l("الصنف", "Catégorie")[lang], domain.category_label(scenario.category)[lang]],
                [_l("معامل الطلب", "Facteur de demande")[lang], p.get("demand_factor")],
                [_l("المدى (ساعات)", "Horizon (heures)")[lang], p.get("horizon_hours")],
                [_l("الطقس", "Météo")[lang], p.get("weather")],
                [_l("البذرة", "Graine")[lang], str(scenario.seed)],
                [_l("عدد المحاكاة", "Tirages")[lang], scenario.runs],
                [_l("الثقة", "Confiance")[lang], result.get("confidence")],
                [_l("المعايرة", "Calibration")[lang], result.get("calibration")],
            ],
        ),
        Section(
            _l("الملخص", "Synthèse"),
            band_cols,
            [
                [_l("بلاغات في الساعة", "Signalements / heure")[lang], *summary["reports_per_hour"].values()],
                [_l("الانتظار (دقيقة)", "Attente (min)")[lang], *summary["wait_minutes"].values()],
                [_l("نسبة الاستعمال", "Utilisation")[lang], *summary["utilisation"].values()],
            ],
        ),
        Section(
            _l("حسب الساعة", "Par heure"),
            [
                Column(_l("الساعة", "Heure")),
                Column(_l("الطلب م50", "Demande P50")),
                Column(_l("الانتظار م50 (دقيقة)", "Attente P50 (min)")),
                Column(_l("الاستعمال م50", "Utilisation P50")),
                Column(_l("تجاوز", "Surcharge")),
            ],
            [
                [
                    datetime.fromisoformat(h["at"]).astimezone(domain.TUNIS_TZ).strftime("%d/%m %H:%M"),
                    h["demand"]["p50"],
                    h["wait_minutes"]["p50"],
                    h["utilisation"]["p50"],
                    "⚠" if h["overloaded"] else "",
                ]
                for h in result["hours"]
            ],
        ),
        Section(
            _l("مصدر المعطيات", "Provenance des données"),
            [Column(_l("البند", "Élément")), Column(_l("القيمة", "Valeur"))],
            [[k, v] for k, v in prov.items()],
        ),
        Section(
            _l("الفرضيات والتحفظات", "Hypothèses et réserves"),
            [Column(_l("النص", "Texte"))],
            [[a[lang]] for a in result["assumptions"]]
            + [[r[lang]] for r in result.get("confidence_reasons", [])]
            + [[result["calibration_note"][lang]], [result["disclaimer"][lang]]],
        ),
    ]
    return sections


async def build(session: AsyncSession, report: Report) -> Content:
    names = await service.gouvernorat_names(session)
    if report.kind in ("period_summary", "gouvernorat_summary"):
        sections = await _period_sections(session, report, names, report.kind == "gouvernorat_summary")
    elif report.kind in domain.INCIDENT_KINDS:
        sections = await _incident_sections(session, report, names, report.kind == "after_action")
    else:
        sections = await _scenario_sections(session, report)
    lang = report.lang
    filters: list[tuple[dict[str, str], str]] = []
    if report.filters.get("category"):
        filters.append((_l("الصنف", "Catégorie"), domain.category_label(report.filters["category"])[lang]))
    if report.filters.get("incident_id"):
        filters.append((_l("التدخل", "Incident"), str(report.filters["incident_id"])))
    if report.filters.get("scenario_id"):
        filters.append((_l("السيناريو", "Scénario"), str(report.filters["scenario_id"])))
    return Content(
        title=report.title,
        kind=report.kind,
        lang=lang,
        period=(report.period_from.isoformat(), report.period_to.isoformat()),
        gouvernorats=[names.get(c, {}).get(lang, str(c)) for c in report.gouvernorats]
        or [_l("كل الولايات (وطني)", "Tous les gouvernorats (national)")[lang]],
        filters=filters,
        min_cell=get_settings().analytics_min_cell,
        generated_at=datetime.now(UTC),
        sections=sections,
    )


def cell_out(value: Any, column: Column, min_cell: int) -> Cell:
    if column.count and isinstance(value, int) and not isinstance(value, bool):
        return domain.floor_cell(value, min_cell)
    if value is None:
        return "—"
    return value if isinstance(value, int | float | str) else str(value)


def _cover_rows(content: Content) -> list[tuple[str, str]]:
    lang = content.lang
    rows = [
        (_l("نوع التقرير", "Type de document")[lang], domain.KIND_LABELS[content.kind][lang]),
        (
            _l("الفترة", "Période")[lang],
            _l(f"من {content.period[0]} إلى {content.period[1]}", f"du {content.period[0]} au {content.period[1]}")[
                lang
            ],
        ),
        (_l("الولايات", "Gouvernorats")[lang], ("، " if lang == "ar" else ", ").join(content.gouvernorats)),
    ]
    rows += [(k[lang], v) for k, v in content.filters]
    rows.append(
        (
            _l("تاريخ الإعداد", "Établi le")[lang],
            content.generated_at.astimezone(domain.TUNIS_TZ).strftime("%Y-%m-%d %H:%M") + " (Africa/Tunis)",
        )
    )
    return rows


# ---------------------------------------------------------------- PDF

_CSS = """
@page { size: A4; margin: 16mm 14mm 18mm 14mm;
  @bottom-center { content: counter(page) " / " counter(pages); font-size: 8pt; color: #666; } }
body { font-family: "Noto Naskh Arabic", "Noto Sans Arabic", "Noto Sans", "DejaVu Sans", sans-serif;
  font-size: 9.5pt; color: #1b1b1b; }
html[dir="rtl"] body { direction: rtl; text-align: right; }
html[dir="ltr"] body { direction: ltr; text-align: left; }
.cover { page-break-after: always; }
.cover h1 { font-size: 20pt; color: #b71c1c; margin: 30mm 0 4mm; }
.cover .brand { font-size: 11pt; color: #555; border-bottom: 3px solid #b71c1c; padding-bottom: 6px; }
.rule { margin-top: 12mm; padding: 8px 10px; border: 1px solid #b71c1c; background: #fff5f5; }
h2 { font-size: 11pt; margin: 14px 0 6px; border-bottom: 1px solid #b71c1c; color: #b71c1c; padding-bottom: 2px; }
table { width: 100%; border-collapse: collapse; margin-top: 4px; page-break-inside: auto; }
tr { page-break-inside: avoid; }
th, td { border: 1px solid #ccc; padding: 3px 5px; vertical-align: top; }
th { background: #f4f4f4; font-weight: bold; }
td.num { text-align: center; direction: ltr; unicode-bidi: isolate; }
.small { font-size: 8pt; color: #555; }
"""


def render_html(content: Content) -> str:
    lang = content.lang
    e = html.escape
    rule = domain.floor_rule(content.min_cell)[lang]
    cover = "".join(f"<tr><th>{e(k)}</th><td>{e(v)}</td></tr>" for k, v in _cover_rows(content))
    body: list[str] = []
    for section in content.sections:
        head = "".join(f"<th>{e(c.label[lang])}</th>" for c in section.columns)
        rows = []
        for row in section.rows:
            cells = []
            for value, column in zip(row, section.columns, strict=False):
                out = cell_out(value, column, content.min_cell)
                numeric = isinstance(out, int | float) or (isinstance(out, str) and out.startswith("<"))
                cells.append(f"<td{' class="num"' if numeric else ''}>{e(str(out))}</td>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        empty = f'<tr><td colspan="{len(section.columns)}">—</td></tr>'
        note = f'<p class="small">{e(section.note[lang])}</p>' if section.note else ""
        body.append(f"<h2>{e(section.title[lang])}</h2><table><tr>{head}</tr>{''.join(rows) or empty}</table>{note}")
    brand = "منصة ء-حماية · الحماية المدنية" if lang == "ar" else "e-Himaya · Protection civile"
    return f"""<!doctype html><html lang="{lang}" dir="{"rtl" if lang == "ar" else "ltr"}"><head><meta charset="utf-8">
<style>{_CSS}</style></head><body>
<div class="cover"><div class="brand">{e(brand)}</div><h1>{e(content.title)}</h1>
<table>{cover}</table><div class="rule">{e(rule)}</div></div>
{"".join(body)}
</body></html>"""


def render_pdf(content: Content) -> bytes:
    from weasyprint import HTML

    pdf: bytes = HTML(string=render_html(content)).write_pdf()
    return pdf


# ---------------------------------------------------------------- XLSX

_INVALID_SHEET = str.maketrans({c: " " for c in "[]:*?/\\"})


def render_xlsx(content: Content) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    lang = content.lang
    rtl = lang == "ar"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="F4F4F4")
    wb = Workbook()
    cover = wb.active
    cover.title = _l("الغلاف", "Couverture")[lang]
    cover.sheet_view.rightToLeft = rtl
    cover.append([content.title])
    cover["A1"].font = Font(bold=True, size=14, color="B71C1C")
    cover.append([])
    for key, value in _cover_rows(content):
        cover.append([key, value])
        cover.cell(row=cover.max_row, column=1).font = bold
    cover.append([])
    cover.append([domain.floor_rule(content.min_cell)[lang]])
    cover.cell(row=cover.max_row, column=1).alignment = Alignment(wrap_text=True)
    cover.column_dimensions["A"].width = 32
    cover.column_dimensions["B"].width = 60
    used: set[str] = {cover.title}
    for index, section in enumerate(content.sections, start=1):
        name = f"{index}. {section.title[lang]}".translate(_INVALID_SHEET)[:31]
        while name in used:
            name = name[:28] + f"~{index}"
        used.add(name)
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = rtl
        ws.append([c.label[lang] for c in section.columns])
        for cell in ws[1]:
            cell.font, cell.fill = bold, head_fill
        for row in section.rows:
            ws.append([cell_out(v, c, content.min_cell) for v, c in zip(row, section.columns, strict=False)])
        if section.note:
            ws.append([])
            ws.append([section.note[lang]])
        for col in range(1, len(section.columns) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 24 if col == 1 else 16
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def render(content: Content, fmt: str) -> bytes:
    return render_pdf(content) if fmt == "pdf" else render_xlsx(content)
