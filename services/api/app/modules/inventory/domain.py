"""Reading and checking the resources register files (phase R1, docs/phase-r/design-resources.md §2.2).

One xlsx template per gouvernorat with three sheets, «الآليات», «المعدات» and «الأعوان», a CSV of one of them, or a
JSON file (R4b): an array of records of one of them, or an object holding the three arrays. The columns of a file that
are not the template's are matched with the register's fields by a mapping the user chose from lists (R4b).
Nothing here touches the database: the caller passes the gouvernorat's units and the catalogue, and the rows come back
resolved (unit id, catalogue key, typed specifications) with the errors that drop a row and the warnings that keep it.
"""

import io
import json
import re
from dataclasses import dataclass, field
from datetime import date
from typing import Any, Literal, cast

from app.modules.programmes.domain import (
    FileError,
    _date,
    _integer,
    _keyed,
    _number,
    _records_csv,
    _rows,
    _sheet,
    _text,
    normalize,
)

Entity = Literal["vehicles", "stock", "agents"]
ENTITIES: tuple[Entity, ...] = ("vehicles", "stock", "agents")
KINDS = ("vehicle", "equipment", "skill")
FAMILIES = ("vehicles", "equipment", "boats", "air", "ppe")
# What the register shows for a vehicle; "assigned" is set by the interventions only.
VEHICLE_STATES = ("ready", "assigned", "maintenance", "broken")
WRITABLE_STATES = ("ready", "maintenance", "broken")
AVAILABILITY = ("on_duty", "rest", "leave", "sick", "training", "mission")
STALE_DAYS = 30
MAX_ROWS = 5000
MAX_BYTES = 10 * 1024 * 1024
# R4b: a JSON file is records of scalars, at most this deep (the object, its arrays, their records, a list of skills).
MAX_JSON_DEPTH = 4
# R4b: how many columns of one sheet a mapping may rename, and how long a heading may be.
MAX_MAPPED = 60
MAX_HEADING = 100
KEY = re.compile(r"^[a-z][a-z0-9_]{1,39}$")

SHEETS: dict[Entity, tuple[str, ...]] = {
    "vehicles": ("الآليات", "vehicles", "véhicules"),
    "stock": ("المعدات", "equipment", "stock", "matériel"),
    "agents": ("الأعوان", "agents", "staff"),
}
STATE_LABELS = {"جاهزة": "ready", "جاهز": "ready", "صيانة": "maintenance", "معطلة": "broken", "معطل": "broken"}
AVAILABILITY_LABELS = {
    "في الخدمة": "on_duty",
    "راحة": "rest",
    "إجازة": "leave",
    "عطلة مرض": "sick",
    "مرض": "sick",
    "تكوين": "training",
    "مأمورية": "mission",
    "مهمة": "mission",
}
YES = {"1", "نعم", "yes", "oui", "true", "x", "✓"}
NO = {"0", "لا", "no", "non", "false", ""}

# The specification columns of the template: the catalogue key, then the Arabic heading.
SPEC_COLUMNS: dict[str, tuple[str, ...]] = {
    "capacity_l": ("السعة (لتر)", "السعة"),
    "flow_lpm": ("الصبيب (لتر في الدقيقة)", "الصبيب"),
    "length_m": ("الطول (متر)", "الطول"),
    "seats": ("عدد الأشخاص",),
    "power_kva": ("القدرة (كيلوفولط أمبير)", "القدرة"),
    "four_wd": ("دفع رباعي",),
}
COLUMNS: dict[Entity, dict[str, tuple[str, ...]]] = {
    "vehicles": {
        "unit": ("الوحدة",),
        "category": ("الصنف",),
        "label": ("التسمية",),
        "code": ("الرقم الداخلي",),
        "state": ("الحالة",),
        "year": ("سنة الصنع",),
        "notes": ("ملاحظات",),
        **SPEC_COLUMNS,
    },
    "stock": {
        "unit": ("الوحدة",),
        "category": ("الصنف",),
        "qty_ready": ("العدد الجاهز", "الجاهز"),
        "qty_broken": ("العدد المعطل", "المعطل"),
        "notes": ("ملاحظات",),
        **SPEC_COLUMNS,
    },
    "agents": {
        "matricule": ("الرقم الإداري", "المعرف"),
        "full_name": ("الاسم واللقب", "الاسم"),
        "rank": ("الرتبة",),
        "unit": ("الوحدة",),
        "availability": ("الجاهزية",),
        "phone": ("الهاتف (مقيّد)", "الهاتف"),
        "skills": ("الاختصاصات",),
        "notes": ("ملاحظات",),
    },
}


@dataclass
class Problem:
    code: str
    entity: str
    line: int | None = None
    value: str = ""

    def out(self) -> dict[str, Any]:
        return {k: v for k, v in self.__dict__.items() if v not in ("", None)}


Format = Literal["xlsx", "csv", "json"]
Mapping = dict[Entity, dict[str, str]]


@dataclass
class Parsed:
    format: Format
    rows: dict[Entity, list[dict[str, Any]]]
    errors: list[Problem] = field(default_factory=list)
    warnings: list[Problem] = field(default_factory=list)
    # R4b: the columns of the file by sheet, and those that match no field of the register (to map them on screen).
    headings: dict[Entity, list[str]] = field(default_factory=dict)
    unknown: dict[Entity, list[str]] = field(default_factory=dict)


@dataclass(frozen=True)
class CategoryRef:
    key: str
    kind: str
    specs: tuple[tuple[str, str], ...]  # (key, number|bool)
    levels: int
    active: bool


# ---------------------------------------------------------------- vocabulary


def vehicle_state(status: str, out_reason: str | None) -> str:
    if status == "available":
        return "ready"
    if status == "assigned":
        return "assigned"
    return out_reason or "broken"


def state_columns(state: str) -> tuple[str, str | None]:
    """The ``ops.resource`` status and out-of-service reason of a register state (ready, maintenance, broken)."""
    if state == "ready":
        return "available", None
    return "out_of_service", state


def check_specs(values: dict[str, Any], category: CategoryRef) -> dict[str, Any]:
    """Only the category's specifications, typed; ``ValueError(key)`` on a value of the wrong type."""
    out: dict[str, Any] = {}
    for key, kind in category.specs:
        value = values.get(key)
        if value in (None, ""):
            continue
        if kind == "bool":
            if isinstance(value, bool):
                out[key] = value
            elif _text(value).lower() in YES:
                out[key] = True
            elif _text(value).lower() in NO:
                out[key] = False
            else:
                raise ValueError(key)
        else:
            try:
                number = _number(value)
            except ValueError as exc:
                raise ValueError(key) from exc
            if number is None or number < 0 or number > 1_000_000:
                raise ValueError(key)
            out[key] = int(number) if number.is_integer() else round(number, 2)
    return out


def matches(specs: dict[str, Any], wanted: dict[str, Any]) -> bool:
    """Whether a vehicle or stock line meets ``{"capacity_l": [2000, None], "four_wd": true}`` (a range or a value)."""
    for key, condition in wanted.items():
        value = specs.get(key)
        if isinstance(condition, list):
            low, high = [*condition, None, None][:2]
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                return False
            if (low is not None and value < low) or (high is not None and value > high):
                return False
        elif value != condition:
            return False
    return True


# ---------------------------------------------------------------- reading


def _format(file_name: str, data: bytes) -> Format:
    if len(data) > MAX_BYTES:
        raise FileError("file_too_large")
    if data[:2] == b"PK" or file_name.lower().endswith(".xlsx"):
        return "xlsx"
    if file_name.lower().endswith(".json") or data.lstrip(b"\xef\xbb\xbf \t\r\n")[:1] in (b"[", b"{"):
        return "json"
    if file_name.lower().endswith((".csv", ".txt")):
        return "csv"
    raise FileError("file_format")


def _records_xlsx(data: bytes) -> dict[Entity, list[dict[str, Any]]]:
    import zipfile

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise FileError("file_unreadable") from exc
    try:
        found: dict[Entity, list[dict[str, Any]]] = {}
        for entity, names in SHEETS.items():
            sheet = _sheet(workbook, *names)
            found[entity] = _rows(sheet) if sheet is not None else []
        if not any(found.values()):
            raise FileError("file_sheet_missing", "الآليات / المعدات / الأعوان")
        return found
    finally:
        workbook.close()


def _depth(value: Any, level: int = 1) -> int:
    if isinstance(value, dict):
        return max((_depth(v, level + 1) for v in value.values()), default=level)
    if isinstance(value, list):
        return max((_depth(v, level + 1) for v in value), default=level)
    return level - 1


def _scalar(value: Any) -> Any:
    """A JSON value of a record: a scalar as it is, a list of scalars (skills) joined; anything else refused."""
    if value is None or isinstance(value, str | int | float | bool):
        return value
    if isinstance(value, list) and all(isinstance(v, str | int | float) for v in value):
        return "؛ ".join(_text(v) for v in value)
    raise FileError("file_json_shape")


def _records_json(data: bytes, entity: Entity | None) -> dict[Entity, list[dict[str, Any]]]:
    try:
        document = json.loads(data.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise FileError("file_encoding") from exc
    except (ValueError, RecursionError) as exc:
        raise FileError("file_unreadable") from exc
    if _depth(document) > MAX_JSON_DEPTH:
        raise FileError("file_json_shape")
    found: dict[Entity, list[Any]] = {e: [] for e in ENTITIES}
    if isinstance(document, list):
        if entity is None:
            raise FileError("entity_required")
        found[entity] = document
    elif isinstance(document, dict):
        names = {normalize(n): e for e, labels in SHEETS.items() for n in (e, *labels)}
        for key, records in document.items():
            target = names.get(normalize(_text(key)))
            if target is None:
                continue
            if not isinstance(records, list):
                raise FileError("file_json_shape")
            found[target] = records
        if not any(found.values()):
            raise FileError("file_sheet_missing", "vehicles / stock / agents")
    else:
        raise FileError("file_json_shape")
    out: dict[Entity, list[dict[str, Any]]] = {}
    for name, records in found.items():
        if not all(isinstance(r, dict) for r in records):
            raise FileError("file_json_shape")
        out[name] = [{_text(k): _scalar(v) for k, v in r.items()} for r in records]
    return out


def check_mapping(mapping: Any) -> Mapping:
    """A mapping chosen on screen: per sheet, a heading of the file to a field of the register. ``FileError`` when a
    sheet or a field is not one of ours, or the mapping is too large: the headings are data and never reach SQL."""
    if not isinstance(mapping, dict):
        raise FileError("mapping_invalid")
    out: Mapping = {}
    for name, pairs in mapping.items():
        if name not in ENTITIES or not isinstance(pairs, dict) or len(pairs) > MAX_MAPPED:
            raise FileError("mapping_invalid", _text(name)[:40])
        entity = cast(Entity, name)
        clean: dict[str, str] = {}
        for heading, target in pairs.items():
            if not isinstance(heading, str) or not isinstance(target, str):
                raise FileError("mapping_invalid")
            heading = _text(heading)
            if not heading or len(heading) > MAX_HEADING or target not in COLUMNS[entity]:
                raise FileError("mapping_invalid", heading[:40] or target[:40])
            clean[heading] = target
        if len(set(clean.values())) != len(clean):
            raise FileError("mapping_duplicate", entity)
        if clean:
            out[entity] = clean
    return out


def _mapped(row: dict[str, Any], pairs: dict[str, str]) -> dict[str, Any]:
    """The row with the mapped headings renamed to the register's field keys (which ``_keyed`` reads first)."""
    if not pairs:
        return row
    by_heading = {normalize(_text(k)): v for k, v in row.items() if k is not None}
    out = dict(row)
    for heading, target in pairs.items():
        value = by_heading.get(normalize(heading))
        if value not in (None, ""):
            out[target] = value
    return out


def _unknown(headings: list[str], entity: Entity, pairs: dict[str, str]) -> list[str]:
    known = {normalize(h) for key, labels in COLUMNS[entity].items() for h in (key, *labels)}
    known |= {normalize(h) for h in pairs}
    return [h for h in headings if h and normalize(h) not in known]


def parse(
    file_name: str,
    data: bytes,
    entity: Entity | None,
    units: dict[str, Any],
    categories: dict[str, CategoryRef],
    mapping: Mapping | None = None,
) -> Parsed:
    """``units`` maps a normalised unit name (Arabic or French) to its id; ``categories`` maps a normalised key or
    name (Arabic, French or English) to the catalogue entry. A CSV, or a JSON array, holds one entity, named by
    ``entity``. ``mapping`` (R4b) renames the columns of a file that are not the template's."""
    fmt = _format(file_name, data)
    if fmt == "xlsx":
        raw = _records_xlsx(data)
    elif fmt == "json":
        raw = _records_json(data, entity)
    else:
        if entity is None:
            raise FileError("entity_required")
        raw = {e: [] for e in ENTITIES}
        raw[entity] = _records_csv(data)
    if sum(len(v) for v in raw.values()) > MAX_ROWS:
        raise FileError("file_too_many_rows")
    mapping = mapping or {}
    parsed = Parsed(format=fmt, rows={e: [] for e in ENTITIES})
    for name, rows in raw.items():
        pairs = mapping.get(name, {})
        if rows:
            headings = list(dict.fromkeys(_text(h) for r in rows[:50] for h in r if h is not None))[:MAX_MAPPED]
            parsed.headings[name] = headings
            parsed.unknown[name] = _unknown(headings, name, pairs)
        for line, row in enumerate(rows, start=2):
            values = _keyed(_mapped(row, pairs), COLUMNS[name])
            if not values:
                continue
            checked = _check(name, values, line, units, categories, parsed)
            if checked is not None:
                parsed.rows[name].append(checked)
    _duplicates(parsed)
    return parsed


def _problem(entity: Entity, code: str, line: int, value: Any = "") -> Problem:
    return Problem(code=code, entity=entity, line=line, value=_text(value))


def _check(
    entity: Entity,
    values: dict[str, Any],
    line: int,
    units: dict[str, Any],
    categories: dict[str, CategoryRef],
    parsed: Parsed,
) -> dict[str, Any] | None:
    errors = parsed.errors
    unit_name = _text(values.get("unit"))
    unit_id = units.get(normalize(unit_name))
    if unit_id is None:
        errors.append(_problem(entity, "unit_unknown", line, unit_name))
        return None
    out: dict[str, Any] = {"line": line, "unit_id": unit_id, "notes": _text(values.get("notes"))[:500] or None}
    if entity == "agents":
        return _check_agent(values, line, out, categories, parsed)
    label = _text(values.get("category"))
    category = categories.get(normalize(label))
    wanted_kind = "vehicle" if entity == "vehicles" else "equipment"
    if category is None or category.kind != wanted_kind:
        errors.append(_problem(entity, "category_unknown", line, label))
        return None
    if not category.active:
        errors.append(_problem(entity, "category_inactive", line, label))
        return None
    out["category_key"] = category.key
    try:
        out["specs"] = check_specs(values, category)
    except ValueError as exc:
        errors.append(_problem(entity, "spec_invalid", line, str(exc)))
        return None
    if entity == "vehicles":
        out["label"] = _text(values.get("label"))[:60]
        if not out["label"]:
            errors.append(_problem(entity, "label_required", line))
            return None
        out["code"] = _text(values.get("code"))[:40] or None
        state = _text(values.get("state")) or "ready"
        state = (
            state
            if state in WRITABLE_STATES
            else {normalize(k): v for k, v in STATE_LABELS.items()}.get(normalize(state), state)
        )
        if state not in WRITABLE_STATES:
            parsed.warnings.append(_problem(entity, "state_unknown", line, state))
            state = "ready"
        out["state"] = state
        try:
            year = _integer(values.get("year"))
        except ValueError:
            year = -1
        if year is not None and not 1950 <= year <= 2100:
            parsed.warnings.append(_problem(entity, "year_invalid", line, values.get("year")))
            year = None
        out["year"] = year
        return out
    try:
        out["qty_ready"] = _integer(values.get("qty_ready")) or 0
        out["qty_broken"] = _integer(values.get("qty_broken")) or 0
    except ValueError as exc:
        errors.append(_problem(entity, "quantity_invalid", line, str(exc)))
        return None
    if out["qty_ready"] > 100_000 or out["qty_broken"] > 100_000:
        errors.append(_problem(entity, "quantity_invalid", line, out["qty_ready"]))
        return None
    if out["qty_ready"] + out["qty_broken"] == 0:
        parsed.warnings.append(_problem(entity, "quantity_zero", line))
    return out


def parse_skills(text: str, categories: dict[str, CategoryRef]) -> list[dict[str, Any]]:
    """«غوص:2:2027-06-30؛ إنقاذ بالحبال» → skills with their level and certificate expiry; ``ValueError(part)``."""
    out: dict[str, dict[str, Any]] = {}
    for part in re.split(r"[;؛,،\n]+", text):
        part = part.strip()
        if not part:
            continue
        pieces = [p.strip() for p in part.split(":")]
        category = categories.get(normalize(pieces[0]))
        if category is None or category.kind != "skill" or not category.active:
            raise ValueError(pieces[0])
        level = 0
        expires: date | None = None
        try:
            if len(pieces) > 1 and pieces[1]:
                level = int(pieces[1])
            if len(pieces) > 2 and pieces[2]:
                expires = _date(pieces[2])
        except ValueError as exc:
            raise ValueError(part) from exc
        if category.levels == 0:
            level = 0
        elif not 1 <= level <= category.levels:
            raise ValueError(part)
        out[category.key] = {"skill_key": category.key, "level": level, "cert_expires": expires}
    return list(out.values())


def _check_agent(
    values: dict[str, Any], line: int, out: dict[str, Any], categories: dict[str, CategoryRef], parsed: Parsed
) -> dict[str, Any] | None:
    matricule = _text(values.get("matricule"))
    name = _text(values.get("full_name"))
    if not 2 <= len(matricule) <= 30:
        parsed.errors.append(_problem("agents", "matricule_required", line, matricule))
        return None
    if not 2 <= len(name) <= 120:
        parsed.errors.append(_problem("agents", "name_required", line, matricule))
        return None
    availability = _text(values.get("availability")) or "on_duty"
    if availability not in AVAILABILITY:
        availability = {normalize(k): v for k, v in AVAILABILITY_LABELS.items()}.get(normalize(availability), "")
    if availability not in AVAILABILITY:
        parsed.warnings.append(_problem("agents", "availability_unknown", line, values.get("availability")))
        availability = "on_duty"
    try:
        skills = parse_skills(_text(values.get("skills")), categories)
    except ValueError as exc:
        parsed.errors.append(_problem("agents", "skill_invalid", line, str(exc)))
        return None
    out.update(
        matricule=matricule,
        full_name=name,
        rank=_text(values.get("rank"))[:60] or None,
        phone=_text(values.get("phone"))[:30] or None,
        availability=availability,
        skills=skills,
    )
    return out


def _duplicates(parsed: Parsed) -> None:
    for entity in ENTITIES:
        seen: set[Any] = set()
        kept = []
        for row in parsed.rows[entity]:
            k = _row_key(entity, row)
            if k in seen:
                parsed.errors.append(_problem(entity, "duplicate_row", row["line"]))
                continue
            seen.add(k)
            kept.append(row)
        parsed.rows[entity] = kept


def _row_key(entity: Entity, row: dict[str, Any]) -> Any:
    """What identifies a row in the register: vehicles by unit and label, equipment by unit, type and specifications,
    staff by matricule."""
    if entity == "vehicles":
        return row["unit_id"], row["label"]
    if entity == "stock":
        return row["unit_id"], row["category_key"], tuple(sorted(row["specs"].items()))
    return row["matricule"]


# ---------------------------------------------------------------- template


def template(units: list[str], categories: list[tuple[str, str, int]]) -> bytes:
    """The xlsx template of a gouvernorat: the three sheets with their headings, and a sheet of the accepted values
    (``categories``: kind, Arabic name, levels)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    first = True
    for entity, names in SHEETS.items():
        sheet = book.active if first else book.create_sheet()
        first = False
        assert sheet is not None
        sheet.title = names[0]
        sheet.sheet_view.rightToLeft = True
        sheet.append([labels[0] for labels in COLUMNS[entity].values()])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    values = book.create_sheet("القيم المقبولة")
    values.sheet_view.rightToLeft = True
    values.append(["الوحدات", "أصناف الآليات", "أصناف المعدات", "الاختصاصات (المستوى الأقصى)", "الحالة", "الجاهزية"])
    for cell in values[1]:
        cell.font = Font(bold=True)
    vehicles = [n for k, n, _ in categories if k == "vehicle"]
    equipment = [n for k, n, _ in categories if k == "equipment"]
    skills = [f"{n} ({lv})" if lv else n for k, n, lv in categories if k == "skill"]
    states = ["جاهزة", "صيانة", "معطلة"]
    availability = ["في الخدمة", "راحة", "إجازة"]
    columns = (units, vehicles, equipment, skills, states, availability)
    for i in range(max(len(c) for c in columns)):
        values.append([c[i] if i < len(c) else None for c in columns])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
