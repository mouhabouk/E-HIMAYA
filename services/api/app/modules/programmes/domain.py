"""Reading and checking seasonal programme files, and comparing two versions (phase L2, design-l2.md §2.5).

Three formats, one model (``data/programmes/2026``): the unified xlsx template (sheets «البرامج» and «النقاط»), a CSV
of points with the same columns as ``programme_points_2026.csv``, and JSON following ``program_schema.json``. Nothing
here touches the database: the gouvernorat names are resolved by the caller.

Errors drop a point (unknown gouvernorat, coordinates outside Tunisia, duplicate reference, missing place); warnings
keep it (no position yet, more lifeguards deployed than planned) and the totals stated in the file are compared with
the sum of the points, as the 2026 beach table does not add up (design §2.5).
"""

import csv
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Literal

Format = Literal["xlsx", "csv", "json"]
Kind = Literal["road", "seasonal", "beach", "other"]
KINDS: tuple[Kind, ...] = ("road", "seasonal", "beach", "other")
STATUSES = ("draft", "approved", "updated", "ended", "cancelled")
POINT_STATUSES = ("open", "closed", "postponed")
CONFIDENCES = ("confirmed", "high", "medium", "approximate", "review", "unknown")
LAT = (30.0, 37.6)
LON = (7.4, 11.7)
CODE = re.compile(r"^[A-Z]+-[0-9]{4}(-[A-Z0-9]+)?$")
MAX_POINTS = 5000
MAX_BYTES = 10 * 1024 * 1024

KIND_LABELS: dict[str, Kind] = {
    "نقاط النجدة والإسعاف بالطرقات": "road",
    "المراكز الموسمية": "seasonal",
    "التركز بالشواطئ": "beach",
    "برنامج آخر": "other",
}
STATUS_LABELS = {"مسودة": "draft", "معتمد": "approved", "مُحيَّن": "updated", "منتهٍ": "ended", "ملغى": "cancelled"}
POINT_STATUS_LABELS = {"مفتوحة": "open", "مغلقة": "closed", "مؤجلة": "postponed"}
# The accuracy of a position starts the label of the 2026 files ("تقريبية: مركز قرية (البلدة لا النقطة)").
CONFIDENCE_LABELS = {
    "مؤكدة": "confirmed",
    "مرتفعة": "high",
    "متوسطة": "medium",
    "تقريبية": "approximate",
    "للمراجعة": "review",
    "غير محدد": "unknown",
}

# Columns of a point: the JSON key, then the headings of the xlsx template.
POINT_COLUMNS: dict[str, tuple[str, ...]] = {
    "program_id": ("معرّف البرنامج",),
    "point_id": ("معرّف النقطة",),
    "governorate": ("الولاية",),
    "place_name": ("المكان / النقطة الرئيسية", "المكان"),
    "sub_points": ("النقاط الفرعية",),
    "latitude": ("خط العرض",),
    "longitude": ("خط الطول",),
    "geo_confidence": ("دقة التموقع",),
    "staff_count": ("عدد الأعوان",),
    "staff_composition": ("تركيبة الفريق",),
    "supervisors": ("عدد المشرفين",),
    "lifeguards_planned": ("السباحون المنقذون (المبرمج)",),
    "lifeguards_deployed": ("السباحون المنقذون (المركّز)",),
    "equipment": ("الإمكانيات المادية",),
    "hours": ("توقيت العمل",),
    "opening_date": ("تاريخ الفتح",),
    "closing_date": ("تاريخ الغلق",),
    "point_status": ("حالة النقطة",),
    "supervisor_name": ("المشرف (مقيّد)", "المشرف"),
    "supervisor_phone": ("هاتف المشرف (مقيّد)", "هاتف المشرف"),
    "notes": ("ملاحظات",),
}
PROGRAMME_COLUMNS: dict[str, tuple[str, ...]] = {
    "program_id": ("معرّف البرنامج",),
    "program_type": ("نوع البرنامج",),
    "title": ("العنوان",),
    "season": ("الموسم",),
    "valid_from": ("ساري من",),
    "valid_to": ("ساري إلى",),
    "version": ("النسخة",),
    "status": ("الحالة",),
    "source_document": ("الوثيقة المصدر",),
    "declared_staff": ("الأعوان",),
    "declared_lifeguards_planned": ("المنقذون المبرمجون",),
    "declared_lifeguards_deployed": ("المنقذون المركّزون",),
}
INTEGERS = ("staff_count", "supervisors", "lifeguards_planned", "lifeguards_deployed")
TEXTS = ("sub_points", "staff_composition", "equipment", "hours", "supervisor_name", "supervisor_phone", "notes")
# Totals compared with the sum of the points: stated key -> point column.
TOTALS = {
    "staff": "staff_count",
    "lifeguards_planned": "lifeguards_planned",
    "lifeguards_deployed": "lifeguards_deployed",
}


class FileError(ValueError):
    """The file cannot be read at all (format, size, structure); the code is the API error code."""

    def __init__(self, code: str, detail: str = "") -> None:
        super().__init__(code)
        self.code = code
        self.detail = detail


@dataclass
class Problem:
    code: str
    programme: str = ""
    ref: str = ""
    line: int | None = None
    value: str = ""

    def out(self) -> dict[str, Any]:
        return {k: v for k, v in self.__dict__.items() if v not in ("", None)}


@dataclass
class ParsedProgramme:
    code: str
    meta: dict[str, Any] | None  # None: the file only has points (CSV), the programme must already exist
    points: list[dict[str, Any]] = field(default_factory=list)
    declared: dict[str, int] = field(default_factory=dict)


@dataclass
class Parsed:
    format: Format
    programmes: list[ParsedProgramme]
    errors: list[Problem] = field(default_factory=list)
    warnings: list[Problem] = field(default_factory=list)


def normalize(value: str) -> str:
    """Compare Arabic names without hamza or ta marbuta variants, tatweel, diacritics or extra spaces."""
    value = unicodedata.normalize("NFKC", value)
    value = re.sub(r"[ً-ْـ]", "", value)
    value = value.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ى", "ي")
    return " ".join(value.split())


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return float(_text(value).replace(",", "."))


def _integer(value: Any) -> int | None:
    number = _number(value)
    if number is None:
        return None
    if not number.is_integer() or number < 0:
        raise ValueError(str(value))
    return int(number)


def confidence_of(label: str, has_position: bool) -> str:
    label = _text(label)
    if label in CONFIDENCES:
        return label
    for prefix, code in CONFIDENCE_LABELS.items():
        if label.startswith(prefix):
            return code
    return "approximate" if has_position else "unknown"


def kind_of(label: str) -> Kind | None:
    label = _text(label)
    if label in KINDS:
        return label
    return KIND_LABELS.get(label)


def status_of(label: str) -> str | None:
    label = _text(label)
    if label in STATUSES:
        return label
    return STATUS_LABELS.get(label)


# ---------------------------------------------------------------- reading


def detect_format(file_name: str, data: bytes) -> Format:
    if len(data) > MAX_BYTES:
        raise FileError("file_too_large")
    name = file_name.lower()
    if data[:2] == b"PK" or name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".json") or data.lstrip(b"\xef\xbb\xbf \r\n\t")[:1] in (b"{", b"["):
        return "json"
    if name.endswith(".csv") or name.endswith(".txt"):
        return "csv"
    raise FileError("file_format")


def _keyed(row: dict[str, Any], columns: dict[str, tuple[str, ...]]) -> dict[str, Any]:
    """A row by JSON key, whether its headings are the JSON keys or the Arabic headings of the template."""
    out: dict[str, Any] = {}
    headings = {normalize(_text(k)): v for k, v in row.items() if k is not None}
    for key, labels in columns.items():
        for heading in (key, *labels):
            found = headings.get(normalize(heading))
            if found not in (None, ""):
                out[key] = found
                break
    return out


def _records_csv(data: bytes) -> list[dict[str, Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileError("file_encoding") from exc
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [dict(r) for r in csv.DictReader(io.StringIO(text), delimiter=delimiter)]


def _sheet(workbook: Any, *names: str) -> Any:
    wanted = {normalize(n) for n in names}
    for sheet in workbook.worksheets:
        if normalize(sheet.title) in wanted:
            return sheet
    return None


def _rows(sheet: Any) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        headings = [_text(h) for h in next(rows)]
    except StopIteration:
        return []
    out = []
    for values in rows:
        if values is None or all(v in (None, "") for v in values):
            continue
        out.append({h: v for h, v in zip(headings, values, strict=False) if h})
    return out


def _records_xlsx(data: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise FileError("file_unreadable") from exc
    try:
        programmes = _sheet(workbook, "البرامج", "programmes")
        points = _sheet(workbook, "النقاط", "points")
        if points is None:
            raise FileError("file_sheet_missing", "النقاط")
        return (_rows(programmes) if programmes is not None else []), _rows(points)
    finally:
        workbook.close()


def _json(data: bytes) -> dict[str, Any]:
    try:
        loaded = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise FileError("file_unreadable") from exc
    if isinstance(loaded, list):
        loaded = {"programmes": loaded}
    if not isinstance(loaded, dict) or not isinstance(loaded.get("programmes"), list):
        raise FileError("file_schema", "programmes")
    return loaded


def parse(file_name: str, data: bytes) -> Parsed:
    fmt = detect_format(file_name, data)
    metas: list[dict[str, Any]] = []
    points: list[dict[str, Any]] = []
    if fmt == "json":
        for p in _json(data)["programmes"]:
            if not isinstance(p, dict):
                raise FileError("file_schema", "programme")
            metas.append({k: v for k, v in p.items() if k != "points"})
            for point in p.get("points") or []:
                if isinstance(point, dict):
                    points.append({"program_id": p.get("program_id"), **point})
    elif fmt == "xlsx":
        raw_metas, raw_points = _records_xlsx(data)
        metas = [_keyed(r, PROGRAMME_COLUMNS) for r in raw_metas]
        points = [_keyed(r, POINT_COLUMNS) for r in raw_points]
    else:
        points = [_keyed(r, POINT_COLUMNS) for r in _records_csv(data)]
    if len(points) > MAX_POINTS:
        raise FileError("file_too_many_points")
    return check(fmt, metas, points)


# ---------------------------------------------------------------- checking


def _meta(raw: dict[str, Any], problems: list[Problem]) -> dict[str, Any] | None:
    code = _text(raw.get("program_id"))
    kind = kind_of(_text(raw.get("program_type")))
    title = _text(raw.get("title"))
    try:
        valid_from = _date(raw.get("valid_from"))
        valid_to = _date(raw.get("valid_to"))
    except ValueError:
        problems.append(Problem("bad_date", programme=code))
        return None
    if not CODE.match(code):
        problems.append(Problem("bad_code", programme=code, value=code))
        return None
    if kind is None:
        problems.append(Problem("unknown_type", programme=code, value=_text(raw.get("program_type"))))
        return None
    if not title or valid_from is None or valid_to is None:
        problems.append(Problem("programme_incomplete", programme=code))
        return None
    if valid_to < valid_from:
        problems.append(Problem("dates_reversed", programme=code))
        return None
    return {
        "code": code,
        "kind": kind,
        "title": title,
        "season": _text(raw.get("season")),
        "valid_from": valid_from,
        "valid_to": valid_to,
        "source_document": _text(raw.get("source_document")),
    }


def _declared(raw: dict[str, Any]) -> dict[str, int]:
    out: dict[str, int] = {}
    for total in TOTALS:
        value = raw.get(f"declared_{total}")
        if isinstance(value, str) and value.startswith("="):
            continue  # a formula of the template: computed, not stated
        try:
            number = _integer(value)
        except ValueError:
            continue
        if number is not None:
            out[total] = number
    return out


def _point(
    line: int, raw: dict[str, Any], code: str, errors: list[Problem], warnings: list[Problem]
) -> dict[str, Any] | None:
    ref = _text(raw.get("point_id"))
    place = _text(raw.get("place_name"))
    if not ref or not place:
        errors.append(Problem("point_incomplete", programme=code, ref=ref, line=line))
        return None
    out: dict[str, Any] = {"ref": ref, "place_name": place, "governorate": _text(raw.get("governorate"))}
    try:
        lat, lon = _number(raw.get("latitude")), _number(raw.get("longitude"))
    except ValueError:
        errors.append(Problem("bad_coordinates", programme=code, ref=ref, line=line))
        return None
    if (lat is None) != (lon is None):
        errors.append(Problem("bad_coordinates", programme=code, ref=ref, line=line))
        return None
    if lat is not None and lon is not None and not (LAT[0] <= lat <= LAT[1] and LON[0] <= lon <= LON[1]):
        errors.append(Problem("outside_tunisia", programme=code, ref=ref, line=line, value=f"{lat},{lon}"))
        return None
    out["lat"], out["lon"] = lat, lon
    note = _text(raw.get("geo_confidence"))
    out["geo_confidence"] = confidence_of(note, lat is not None)
    out["geo_note"] = note or None
    for key in INTEGERS:
        try:
            out[key] = _integer(raw.get(key))
        except ValueError:
            errors.append(Problem("bad_number", programme=code, ref=ref, line=line, value=key))
            return None
    for key in TEXTS:
        out[key] = _text(raw.get(key)) or None
    try:
        out["opening_date"] = _date(raw.get("opening_date"))
        out["closing_date"] = _date(raw.get("closing_date"))
    except ValueError:
        errors.append(Problem("bad_date", programme=code, ref=ref, line=line))
        return None
    status = _text(raw.get("point_status"))
    out["status"] = status if status in POINT_STATUSES else POINT_STATUS_LABELS.get(status, "open")
    planned, deployed = out["lifeguards_planned"], out["lifeguards_deployed"]
    if planned is not None and deployed is not None and deployed > planned:
        warnings.append(Problem("deployed_above_planned", programme=code, ref=ref, line=line))
    return out


def totals(points: list[dict[str, Any]]) -> dict[str, int]:
    return {total: sum(p.get(column) or 0 for p in points) for total, column in TOTALS.items()}


def check(fmt: Format, metas: list[dict[str, Any]], raw_points: list[dict[str, Any]]) -> Parsed:
    errors: list[Problem] = []
    warnings: list[Problem] = []
    found: dict[str, ParsedProgramme] = {}
    for raw in metas:
        meta = _meta(raw, errors)
        if meta is None:
            continue
        if meta["code"] in found:
            errors.append(Problem("duplicate_programme", programme=meta["code"]))
            continue
        found[meta["code"]] = ParsedProgramme(meta["code"], meta, declared=_declared(raw))
    seen: set[tuple[str, str]] = set()
    for line, raw in enumerate(raw_points, start=2):
        code = _text(raw.get("program_id"))
        if not code:
            errors.append(Problem("programme_missing", line=line, ref=_text(raw.get("point_id"))))
            continue
        target = found.get(code)
        if target is None:
            if metas:  # a full file (xlsx or JSON) must describe every programme it has points for
                errors.append(Problem("programme_unknown", programme=code, line=line))
                continue
            target = found[code] = ParsedProgramme(code, None)
        point = _point(line, raw, code, errors, warnings)
        if point is None:
            continue
        if (code, point["ref"]) in seen:
            errors.append(Problem("duplicate_point", programme=code, ref=point["ref"], line=line))
            continue
        seen.add((code, point["ref"]))
        if point["lat"] is None:
            warnings.append(Problem("no_position", programme=code, ref=point["ref"], line=line))
        target.points.append(point)
    for programme in found.values():
        computed = totals(programme.points)
        for total, stated in programme.declared.items():
            if stated != computed[total]:
                warnings.append(
                    Problem(f"total_{total}", programme=programme.code, value=f"{stated}≠{computed[total]}")
                )
    return Parsed(fmt, list(found.values()), errors, warnings)


# ---------------------------------------------------------------- versions

PROGRAMME_FIELDS = ("code", "kind", "title", "season", "valid_from", "valid_to", "status", "source_document")
POINT_FIELDS = (
    "gouvernorat_code",
    "place_name",
    "sub_points",
    "lat",
    "lon",
    "geo_confidence",
    "geo_note",
    "staff_count",
    "staff_composition",
    "supervisors",
    "lifeguards_planned",
    "lifeguards_deployed",
    "equipment",
    "hours",
    "opening_date",
    "closing_date",
    "status",
    "supervisor_name",
    "supervisor_phone",
    "notes",
)


def diff(before: dict[str, Any], after: dict[str, Any]) -> dict[str, Any]:
    """What changed between two snapshots: programme fields, points added, removed and changed (field by field)."""
    old_p, new_p = before.get("programme", {}), after.get("programme", {})
    programme = {k: [old_p.get(k), new_p.get(k)] for k in PROGRAMME_FIELDS if old_p.get(k) != new_p.get(k)}
    old_points, new_points = before.get("points", {}), after.get("points", {})
    added = sorted(set(new_points) - set(old_points))
    removed = sorted(set(old_points) - set(new_points))
    changed: dict[str, dict[str, list[Any]]] = {}
    for ref in sorted(set(old_points) & set(new_points)):
        fields = {
            k: [old_points[ref].get(k), new_points[ref].get(k)]
            for k in POINT_FIELDS
            if old_points[ref].get(k) != new_points[ref].get(k)
        }
        if fields:
            changed[ref] = fields
    return {"programme": programme, "added": added, "removed": removed, "changed": changed}
