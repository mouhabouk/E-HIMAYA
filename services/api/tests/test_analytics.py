"""Phase 8 §1: analytics and reports — scope filtering, overview, seasonal trends, hexagon hotspots, what-if scenarios
(seeded determinism, provenance, confidence), reports repository (PDF/XLSX actually produced, k-anonymity floor,
302 download audited), scheduled reports firing with in-app notifications.

Every test seeds its own reports in a past period and a gouvernorat no other test module writes to, so the shared
test database does not disturb the counts.
"""

import hashlib
import io
import random
import uuid
from collections.abc import Callable
from datetime import UTC, date, datetime, time, timedelta
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.db import get_sessionmaker
from app.modules.analytics import domain, scenario, worker
from app.modules.reports.live import _route
from tests.conftest import STAFF_ISS, FakeStore, TokenFactory, bearer, outbox_events

Staff = Callable[..., dict[str, str]]
TZ = domain.TUNIS_TZ
# Gouvernorats reserved to this module's data (other tests write to Sfax 61 and Tunis 11).
SILIANA, KEF, TOZEUR, KEBILI, TATAOUINE, GAFSA = 34, 33, 72, 73, 83, 71
COORDS = {
    SILIANA: (36.08, 9.37),
    KEF: (36.17, 8.71),
    TOZEUR: (33.92, 8.13),
    KEBILI: (33.70, 8.97),
    TATAOUINE: (32.93, 10.45),
    GAFSA: (34.43, 8.78),
}


# ---------------------------------------------------------------- seeding helpers


async def _reset(session: AsyncSession, *codes: int) -> None:
    """Remove what earlier runs seeded in this module's gouvernorats (reports, incidents, fleet)."""
    params = {"c": list(codes)}
    await session.execute(
        text(
            "DELETE FROM ops.assignment WHERE incident_id IN "
            "(SELECT id FROM ops.incident WHERE gouvernorat_code = ANY(:c))"
        ),
        params,
    )
    await session.execute(text("DELETE FROM reports.report WHERE gouvernorat_code = ANY(:c)"), params)
    await session.execute(text("DELETE FROM ops.incident WHERE gouvernorat_code = ANY(:c)"), params)
    await session.execute(
        text("DELETE FROM ops.resource WHERE unit_id IN (SELECT id FROM ops.unit WHERE gouvernorat_code = ANY(:c))"),
        params,
    )
    await session.execute(text("DELETE FROM ops.unit WHERE gouvernorat_code = ANY(:c)"), params)


async def _account(session: AsyncSession) -> uuid.UUID:
    user_id = uuid.uuid4()
    await session.execute(
        text("INSERT INTO identity.account (user_id, kind, keycloak_sub) VALUES (:u, 'citizen', :s)"),
        {"u": user_id, "s": f"analytics-test-{user_id}"},
    )
    return user_id


async def seed_report(
    session: AsyncSession,
    user_id: uuid.UUID,
    at: datetime,
    code: int,
    category: str = "fire",
    status: str = "received",
    source: str = "app",
    incident_id: uuid.UUID | None = None,
    lat: float | None = None,
    lon: float | None = None,
) -> uuid.UUID:
    rid = uuid.uuid4()
    base_lat, base_lon = COORDS[code]
    tracking = f"EH-{random.randint(10, 99)}-{random.randint(10, 99)}-{random.randint(0, 9_999_999):07d}"
    await session.execute(
        text(
            "INSERT INTO reports.report (id, tracking_code, user_id, category, location, gouvernorat_code, status, "
            "source, created_at, received_at, incident_id) VALUES (:id, :tc, :u, :c, "
            "CAST(ST_SetSRID(ST_MakePoint(:lon, :lat), 4326) AS geography), :g, :st, :src, :at, :at, :inc)"
        ),
        {
            "id": rid,
            "tc": tracking,
            "u": user_id,
            "c": category,
            "lat": lat if lat is not None else base_lat,
            "lon": lon if lon is not None else base_lon,
            "g": code,
            "st": status,
            "src": source,
            "at": at,
            "inc": incident_id,
        },
    )
    await session.execute(
        text("INSERT INTO reports.status_history (report_id, status, at, actor) VALUES (:r, 'received', :at, 'test')"),
        {"r": rid, "at": at},
    )
    return rid


async def seed_incident(session: AsyncSession, code: int, category: str, opened: datetime) -> uuid.UUID:
    iid = uuid.uuid4()
    await session.execute(
        text(
            "INSERT INTO ops.incident (id, category, gouvernorat_code, opened_at, status) "
            "VALUES (:id, :c, :g, :at, 'open')"
        ),
        {"id": iid, "c": category, "g": code, "at": opened},
    )
    return iid


async def seed_fleet(session: AsyncSession, code: int, kinds: list[str]) -> list[uuid.UUID]:
    unit = uuid.uuid4()
    lat, lon = COORDS[code]
    await session.execute(
        text(
            "INSERT INTO ops.unit (id, name_ar, name_fr, kind, gouvernorat_code, location) VALUES "
            "(:id, 'وحدة اختبار', 'Unité test', 'unit', :g, "
            "CAST(ST_SetSRID(ST_MakePoint(:lon, :lat), 4326) AS geography))"
        ),
        {"id": unit, "g": code, "lat": lat, "lon": lon},
    )
    ids = []
    for i, kind in enumerate(kinds):
        rid = uuid.uuid4()
        await session.execute(
            text("INSERT INTO ops.resource (id, unit_id, kind, label) VALUES (:id, :u, :k, :l)"),
            {"id": rid, "u": unit, "k": kind, "l": f"T-{i}"},
        )
        ids.append(rid)
    return ids


def local(day: date, hour: int = 12, minute: int = 0) -> datetime:
    return datetime.combine(day, time(hour, minute), tzinfo=TZ)


def national(staff: Staff, roles: list[str] | None = None) -> dict[str, str]:
    return staff(roles or ["central_committee"], gouvernorat=None)


# ---------------------------------------------------------------- pure functions


def test_k_anonymity_floor() -> None:
    assert domain.floor_cell(0, 5) == 0
    assert domain.floor_cell(1, 5) == "<5"
    assert domain.floor_cell(4, 5) == "<5"
    assert domain.floor_cell(5, 5) == 5
    assert domain.floor_cell(None, 5) is None
    assert "<5" in domain.floor_rule(5)["fr"] and "5" in domain.floor_rule(5)["ar"]


def test_schedule_times_in_tunis() -> None:
    after = datetime(2026, 10, 20, 7, 0, tzinfo=UTC)  # 08:00 in Tunis (UTC+1)
    daily = domain.next_run("daily", time(7, 30), None, None, after)
    assert daily.astimezone(TZ) == datetime(2026, 10, 21, 7, 30, tzinfo=TZ)
    assert domain.next_run("daily", time(9, 0), None, None, after).astimezone(TZ) == datetime(
        2026, 10, 20, 9, 0, tzinfo=TZ
    )
    weekly = domain.next_run("weekly", time(6, 0), 1, None, after)  # 2026-10-20 is a Tuesday
    assert weekly.astimezone(TZ) == datetime(2026, 10, 26, 6, 0, tzinfo=TZ)
    monthly = domain.next_run("monthly", time(6, 0), None, 1, after)
    assert monthly.astimezone(TZ) == datetime(2026, 11, 1, 6, 0, tzinfo=TZ)
    december = domain.next_run("monthly", time(6, 0), None, 28, datetime(2026, 12, 28, 6, 0, tzinfo=UTC))
    assert december.astimezone(TZ) == datetime(2027, 1, 28, 6, 0, tzinfo=TZ)
    assert domain.scheduled_period("daily", date(2026, 10, 20)) == (date(2026, 10, 19), date(2026, 10, 19))
    assert domain.scheduled_period("weekly", date(2026, 10, 20)) == (date(2026, 10, 13), date(2026, 10, 19))
    assert domain.scheduled_period("monthly", date(2026, 3, 1)) == (date(2026, 2, 1), date(2026, 2, 28))


def test_simulation_is_deterministic_and_bounded() -> None:
    rates = [2.0] * 12 + [8.0] * 12
    a = scenario.simulate(rates, 0.8, [], 60.0, 3, 200, seed=7)
    b = scenario.simulate(rates, 0.8, [], 60.0, 3, 200, seed=7)
    c = scenario.simulate(rates, 0.8, [], 60.0, 3, 200, seed=8)
    assert a == b
    assert a.run_total != c.run_total
    assert len(a.run_total) == 200 and len(a.reports) == 24 and len(a.reports[0]) == 200
    assert all(0.0 <= u <= 1.0 for hour in a.utilisation for u in hour)
    # Demand above capacity in the second half: waits grow there.
    late = scenario.percentile(a.waits[23], 0.5)
    early = scenario.percentile(a.waits[2], 0.5)
    assert late > early
    mean_reports = sum(a.run_total) / len(a.run_total)
    assert 0.85 * sum(rates) < mean_reports < 1.15 * sum(rates)


def test_live_route_reaches_recipients_only() -> None:
    from app.core.security import Principal, Realm, Role
    from app.modules.reports.live import LiveClient

    route = _route("analytics.report_ready", {"report_id": "x", "recipients": ["sub-a"]})
    assert route is not None
    audience, visible = route

    def client(subject: str) -> Any:
        principal = Principal(subject, Realm.STAFF, frozenset({Role.OPERATOR}), None)
        return LiveClient(None, principal, 0)  # type: ignore[arg-type]

    assert audience == "staff"
    assert visible(client("sub-a")) and not visible(client("sub-b"))


# ---------------------------------------------------------------- overview and scope


@pytest.fixture
async def march_2019(session: AsyncSession) -> dict[str, Any]:
    """Siliana: 3 fire reports and 1 flood, Kef: 2 floods, in March 2019; one Siliana report answered."""
    await _reset(session, SILIANA, KEF)
    user = await _account(session)
    day = date(2019, 3, 12)
    incident = await seed_incident(session, SILIANA, "fire", local(day, 10))
    first = await seed_report(session, user, local(day, 10), SILIANA, "fire", "closed", "app", incident)
    await seed_report(session, user, local(day, 11), SILIANA, "fire", "in_progress", "sms")
    await seed_report(session, user, local(date(2019, 3, 13), 23), SILIANA, "fire", "rejected", "call")
    await seed_report(session, user, local(date(2019, 3, 14), 9), SILIANA, "flood", "received", "portal")
    await seed_report(session, user, local(day, 15), KEF, "flood", "received", "app")
    await seed_report(session, user, local(day, 16), KEF, "flood", "received", "app")
    t0 = local(day, 10)
    await session.execute(
        text(
            "INSERT INTO ops.assignment (incident_id, kind, status, created_by, created_at, on_scene_at) "
            "SELECT :i, 'volunteer', 'on_scene', 'test', :a, :o"
        ),
        {"i": incident, "a": t0 + timedelta(minutes=10), "o": t0 + timedelta(minutes=30)},
    )
    await session.execute(
        text("INSERT INTO reports.status_history (report_id, status, at, actor) VALUES (:r, 'closed', :at, 'test')"),
        {"r": first, "at": t0 + timedelta(hours=2)},
    )
    await session.commit()
    return {"incident": incident, "params": {"from": "2019-03-01", "to": "2019-03-31"}}


async def test_overview_totals_breakdowns_and_response_times(
    client: httpx.AsyncClient, staff: Staff, march_2019: dict[str, Any]
) -> None:
    r = await client.get(
        "/api/v1/staff/analytics/overview",
        params=march_2019["params"] | {"gouvernorat": SILIANA},
        headers=national(staff),
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["from"] == "2019-03-01" and body["gouvernorats"] == [SILIANA]
    assert body["totals"] == {"received": 4, "verified": 2, "in_progress": 1, "closed": 1, "rejected": 1}
    assert {g["key"]: g["count"] for g in body["by_category"]} == {"fire": 3, "flood": 1}
    assert {g["key"]: g["count"] for g in body["by_source"]} == {"app": 1, "sms": 1, "call": 1, "portal": 1}
    assert {g["key"]: g["count"] for g in body["by_hour"]} == {9: 1, 10: 1, 11: 1, 23: 1}  # Africa/Tunis hours
    assert {g["key"]: g["count"] for g in body["by_weekday"]} == {2: 2, 3: 1, 4: 1}
    rt = body["response_time"]
    assert rt["to_assignment"] == {"mean_s": 600.0, "p90_s": 600.0, "n": 1}
    assert rt["to_on_scene"]["mean_s"] == 1800.0
    assert rt["to_closed"]["mean_s"] == 7200.0
    assert body["min_cell"] == 5
    # Category filter.
    r = await client.get(
        "/api/v1/staff/analytics/overview",
        params=march_2019["params"] | {"gouvernorat": SILIANA, "category": "flood"},
        headers=national(staff),
    )
    assert r.json()["totals"]["received"] == 1


async def test_overview_scope_filtering(client: httpx.AsyncClient, staff: Staff, march_2019: dict[str, Any]) -> None:
    params = march_2019["params"]
    # A regional reader sees their own gouvernorat by default, never the national figures.
    r = await client.get("/api/v1/staff/analytics/overview", params=params, headers=staff(["operator"], SILIANA))
    assert r.status_code == 200
    assert r.json()["gouvernorats"] == [SILIANA] and r.json()["totals"]["received"] == 4
    r = await client.get(
        "/api/v1/staff/analytics/overview", params=params | {"gouvernorat": KEF}, headers=staff(["operator"], SILIANA)
    )
    assert r.status_code == 403 and r.json() == {"detail": "out_of_scope"}
    for path in ("trends", "hotspots"):
        r = await client.get(
            f"/api/v1/staff/analytics/{path}", params={"gouvernorat": KEF}, headers=staff(["operator"], SILIANA)
        )
        assert r.status_code == 403, path
    # National: both gouvernorats.
    r = await client.get("/api/v1/staff/analytics/overview", params=params, headers=national(staff))
    by_gov = {g["key"]: g["count"] for g in r.json()["by_gouvernorat"]}
    assert by_gov[SILIANA] == 4 and by_gov[KEF] == 2
    # Roles: citizens, auditors and system administrators read no analytics.
    r = await client.get("/api/v1/staff/analytics/overview", headers=staff(["org_auditor"], None))
    assert r.status_code == 403
    r = await client.get(
        "/api/v1/staff/analytics/overview", params={"from": "2019-03-10", "to": "2019-03-01"}, headers=national(staff)
    )
    assert r.status_code == 422 and r.json() == {"detail": "invalid_period"}


# ---------------------------------------------------------------- trends


async def test_trends_seasonal_baseline(client: httpx.AsyncClient, staff: Staff, session: AsyncSession) -> None:
    await _reset(session, TOZEUR)
    user = await _account(session)
    # Tozeur, January: 2 reports in 2016, 4 in 2017, 6 in 2018, 5 in 2019.
    for year, n in ((2016, 2), (2017, 4), (2018, 6), (2019, 5)):
        for i in range(n):
            await seed_report(session, user, local(date(year, 1, 10 + i)), TOZEUR, "fire")
    await session.commit()
    r = await client.get(
        "/api/v1/staff/analytics/trends",
        params={
            "metric": "reports",
            "interval": "month",
            "from": "2019-01-01",
            "to": "2019-02-28",
            "gouvernorat": TOZEUR,
        },
        headers=national(staff),
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["unit"] == "count" and body["baseline_available"] is True
    jan, feb = body["points"]
    assert jan["period_start"] == "2019-01-01" and jan["value"] == 5.0 and jan["last_year"] == 6.0
    # Same month of 2016, 2017, 2018 (2014-2015 predate the first data): mean 4, sd 2 → 4 ± 3.92.
    assert jan["expected"] == 4.0 and jan["low"] == 0.08 and jan["high"] == 7.92
    assert feb["value"] == 0.0 and feb["expected"] == 0.0
    # Too little history: no band.
    r = await client.get(
        "/api/v1/staff/analytics/trends",
        params={
            "metric": "reports",
            "interval": "month",
            "from": "2017-01-01",
            "to": "2017-01-31",
            "gouvernorat": TOZEUR,
        },
        headers=national(staff),
    )
    point = r.json()["points"][0]
    assert point["value"] == 4.0 and point["last_year"] == 2.0 and point["expected"] is None
    assert r.json()["baseline_available"] is False
    # Weekly and daily series and the other metrics answer with the same shape.
    for metric in ("response_time", "interventions"):
        r = await client.get(
            "/api/v1/staff/analytics/trends",
            params={
                "metric": metric,
                "interval": "week",
                "from": "2019-01-01",
                "to": "2019-01-31",
                "gouvernorat": TOZEUR,
            },
            headers=national(staff),
        )
        assert r.status_code == 200, r.text
        assert r.json()["points"][0]["period_start"] == "2018-12-31"  # ISO week starts on Monday
    r = await client.get(
        "/api/v1/staff/analytics/trends",
        params={"interval": "day", "from": "2010-01-01", "to": "2019-01-01"},
        headers=national(staff),
    )
    assert r.status_code == 422 and r.json()["detail"] == "period_too_long"


# ---------------------------------------------------------------- hotspots


async def test_hotspots_are_hexagons_with_counts(
    client: httpx.AsyncClient, staff: Staff, session: AsyncSession
) -> None:
    await _reset(session, KEBILI)
    user = await _account(session)
    day = date(2019, 5, 2)
    for dlat in (0.0, 0.001, 0.002):
        await seed_report(session, user, local(day), KEBILI, "fire", lat=33.70 + dlat, lon=8.97)
    await seed_report(session, user, local(day), KEBILI, "fire", lat=33.40, lon=8.50)
    await session.commit()
    r = await client.get(
        "/api/v1/staff/analytics/hotspots",
        params={"from": "2019-05-01", "to": "2019-05-31", "gouvernorat": KEBILI, "size_km": 5},
        headers=staff(["operator"], KEBILI),
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["type"] == "FeatureCollection" and body["cell_size_m"] == 5000
    assert [f["properties"]["count"] for f in body["features"]] == [3, 1]
    for feature in body["features"]:
        assert feature["geometry"]["type"] == "Polygon"
        assert len(feature["geometry"]["coordinates"][0]) == 7  # a closed hexagon
    assert "Point" not in r.text


# ---------------------------------------------------------------- what-if scenarios


@pytest.fixture
async def tataouine(session: AsyncSession) -> dict[str, Any]:
    """Tataouine: 40 fire reports on 8 days of the same season last year, 3 fire engines and a water tanker."""
    await _reset(session, TATAOUINE)
    user = await _account(session)
    today = domain.local_today()
    for d in range(8):
        day = today - timedelta(days=365 + d * 3)
        for i in range(5):
            await seed_report(session, user, local(day, (i * 5) % 24, 7), TATAOUINE, "fire", "closed")
    fleet = await seed_fleet(session, TATAOUINE, ["fire_engine", "fire_engine", "fire_engine", "water_tanker"])
    await session.commit()
    return {"fleet": fleet}


def _scenario(**overrides: Any) -> dict[str, Any]:
    body = {
        "name": "Été sec",
        "gouvernorats": [TATAOUINE],
        "category": "fire",
        "demand_factor": 2.0,
        "unavailable_resources": [],
        "horizon_hours": 12,
        "weather": "fire_danger_extreme",
        "seed": 42,
    }
    return body | overrides


async def test_scenario_seeded_determinism_and_provenance(
    client: httpx.AsyncClient, staff: Staff, tataouine: dict[str, Any], session: AsyncSession
) -> None:
    headers = national(staff, ["ops_manager"])
    r1 = await client.post("/api/v1/staff/analytics/scenarios", json=_scenario(), headers=headers)
    assert r1.status_code == 201, r1.text
    r2 = await client.post("/api/v1/staff/analytics/scenarios", json=_scenario(), headers=headers)
    one, two = r1.json(), r2.json()
    assert one["seed"] == 42 and one["runs"] >= 200 and one["status"] == "done"
    res1, res2 = one["result"], two["result"]
    if res1["start_at"] == res2["start_at"]:  # same hour (not across an hour boundary)
        assert res1 == res2
    assert res1["summary"] == res2["summary"] or res1["start_at"] != res2["start_at"]
    other = await client.post("/api/v1/staff/analytics/scenarios", json=_scenario(seed=43), headers=headers)
    assert other.json()["result"]["hours"] != res1["hours"]
    # Shape, rounding and percentiles.
    assert len(res1["hours"]) == 12
    hour = res1["hours"][0]
    assert set(hour["demand"]) == {"p10", "p50", "p90"}
    assert hour["demand"]["p10"] <= hour["demand"]["p50"] <= hour["demand"]["p90"]
    assert all(isinstance(v, int) for v in hour["wait_minutes"].values())
    assert all(0 <= v <= 1 for v in hour["utilisation"].values())
    assert res1["summary"]["capacity"] == 4
    # Provenance, calibration and confidence (audit finding T26).
    prov = res1["provenance"]
    # Covered days: the season windows minus the days without any report nationwide (other tests' data counts).
    assert prov["events"] == 40 and prov["days_covered"] >= 8 and isinstance(prov["gap_days"], int)
    assert prov["history_from"] and prov["history_to"]
    assert prov["weather_factor"] == 2.5 and prov["weather_factor_source"] == "default"
    assert prov["service_time_source"] == "default"
    assert res1["calibration"] == "uncalibrated" and set(res1["calibration_note"]) == {"ar", "fr", "en"}
    assert res1["confidence"] == "low"  # default engagement times (and few days of history)
    assert any("Durées" in reason["fr"] for reason in res1["confidence_reasons"])
    if prov["days_covered"] < scenario.MIN_DAYS:
        assert any("jours" in reason["fr"] for reason in res1["confidence_reasons"])
    assert all(set(a) == {"ar", "fr", "en"} for a in res1["assumptions"]) and len(res1["assumptions"]) >= 5
    assert "historique" in res1["disclaimer"]["fr"]
    # Stored, listed (without the hourly detail), readable, audited.
    r = await client.get(f"/api/v1/staff/analytics/scenarios/{one['id']}", headers=headers)
    assert r.status_code == 200 and r.json()["result"] == res1
    r = await client.get("/api/v1/staff/analytics/scenarios", params={"limit": 2}, headers=headers)
    assert r.status_code == 200 and len(r.json()["items"]) == 2 and r.json()["next_cursor"]
    assert "hours" not in r.json()["items"][0]["result"]
    action = await session.scalar(
        text("SELECT action FROM audit.log WHERE target = :t"), {"t": f"analytics_scenario:{one['id']}"}
    )
    assert action == "analytics.scenario.run"


async def test_scenario_low_confidence_overload_and_capacity(
    client: httpx.AsyncClient, staff: Staff, tataouine: dict[str, Any]
) -> None:
    headers = national(staff, ["central_committee"])
    # Few events: another category with no history in Tataouine.
    r = await client.post(
        "/api/v1/staff/analytics/scenarios",
        json=_scenario(category="chemical_leak", weather="heavy_rain"),
        headers=headers,
    )
    assert r.status_code == 201, r.text
    result = r.json()["result"]
    assert result["confidence"] == "low" and result["provenance"]["events"] == 0
    assert any("< 30" in reason["fr"] for reason in result["confidence_reasons"])
    assert result["provenance"]["weather_factor_source"] == "none"
    # One engine left and ten times the demand: hours over capacity.
    r = await client.post(
        "/api/v1/staff/analytics/scenarios",
        json=_scenario(demand_factor=10.0, unavailable_resources={"fire_engine": 2, "water_tanker": 1}),
        headers=headers,
    )
    assert r.status_code == 201, r.text
    result = r.json()["result"]
    assert result["summary"]["capacity"] == 1 and result["provenance"]["resources_removed"] == 3
    assert result["overload_hours"] and all(result["hours"][h]["overloaded"] for h in result["overload_hours"])
    # Resources by id: unknown ids are refused; removing every resource leaves no capacity.
    r = await client.post(
        "/api/v1/staff/analytics/scenarios",
        json=_scenario(unavailable_resources=[str(uuid.uuid4())]),
        headers=headers,
    )
    assert r.status_code == 422 and r.json() == {"detail": "unknown_resource"}
    r = await client.post(
        "/api/v1/staff/analytics/scenarios",
        json=_scenario(unavailable_resources=[str(i) for i in tataouine["fleet"]]),
        headers=headers,
    )
    assert r.status_code == 422 and r.json() == {"detail": "no_capacity"}
    # Scope and roles.
    r = await client.post("/api/v1/staff/analytics/scenarios", json=_scenario(), headers=staff(["ops_manager"], 61))
    assert r.status_code == 403 and r.json() == {"detail": "out_of_scope"}
    r = await client.post("/api/v1/staff/analytics/scenarios", json=_scenario(), headers=staff(["operator"], None))
    assert r.status_code == 403
    r = await client.post("/api/v1/staff/analytics/scenarios", json=_scenario(horizon_hours=100), headers=headers)
    assert r.status_code == 422


# ---------------------------------------------------------------- reports repository


async def _generate_all() -> None:
    sessionmaker = get_sessionmaker()
    while await worker.generate_next(sessionmaker):
        pass


async def _create(client: httpx.AsyncClient, headers: dict[str, str], **overrides: Any) -> dict[str, Any]:
    body = {
        "kind": "period_summary",
        "title": "Synthèse mars 2019",
        "period_from": "2019-03-01",
        "period_to": "2019-03-31",
        "gouvernorats": [SILIANA],
        "format": "pdf",
        "lang": "ar",
    } | overrides
    r = await client.post("/api/v1/staff/reports-library", json=body, headers=headers)
    assert r.status_code == 202, r.text
    assert r.json()["status"] == "queued"
    return dict(r.json())


async def test_report_pdf_generated_stored_and_downloaded(
    client: httpx.AsyncClient,
    staff: Staff,
    make_token: TokenFactory,
    march_2019: dict[str, Any],
    fake_store: FakeStore,
    session: AsyncSession,
) -> None:
    subject = str(uuid.uuid4())
    headers = bearer(make_token(issuer=STAFF_ISS, roles=["ops_manager"], scope_level="national", sub=subject))
    created = await _create(client, headers)
    await _generate_all()
    r = await client.get(f"/api/v1/staff/reports-library/{created['id']}", headers=headers)
    report = r.json()
    assert report["status"] == "ready", report
    key = f"analytics/reports/{created['id']}.pdf"
    data = fake_store.objects[key]
    assert data.startswith(b"%PDF") and len(data) > 1000
    assert report["sha256"] == hashlib.sha256(data).hexdigest() and report["bytes"] == len(data)
    kind = await session.scalar(text("SELECT kind FROM reports.media WHERE id = :m"), {"m": report["media_id"]})
    assert kind == "analytics_report"
    # The creator is notified in the app and on the live feed.
    events = [e for e in await outbox_events(session, "analytics.report_ready") if e["report_id"] == created["id"]]
    assert events and events[0]["recipients"] == [subject]
    r = await client.get("/api/v1/staff/analytics/notifications", headers=headers)
    assert r.json()["unread"] == 1 and r.json()["items"][0]["report_id"] == created["id"]
    # 302 to a pre-signed URL (or the URL as JSON), audited.
    r = await client.get(f"/api/v1/staff/reports-library/{created['id']}/download", headers=headers)
    assert r.status_code == 302 and r.headers["location"].startswith(f"http://s3.test/{key}")
    assert fake_store.presigned[-1][1] == "application/pdf" and "attachment" in (fake_store.presigned[-1][2] or "")
    r = await client.get(
        f"/api/v1/staff/reports-library/{created['id']}/download", headers=headers | {"Accept": "application/json"}
    )
    assert r.status_code == 200 and r.json()["url"].startswith("http://s3.test/")
    downloads = await session.scalar(
        text("SELECT count(*) FROM audit.log WHERE action = 'analytics.report.download' AND target = :t"),
        {"t": f"analytics_report:{created['id']}"},
    )
    assert downloads == 2
    # HTML of the Arabic template: right to left, cover with the filters and the floor rule, floored counts.
    from app.modules.analytics import documents
    from app.modules.analytics.models import Report

    async with get_sessionmaker()() as s:
        row = await s.get(Report, uuid.UUID(created["id"]))
        assert row is not None
        content = await documents.build(s, row)
    page = documents.render_html(content)
    assert 'dir="rtl"' in page and "&lt;5" in page and "Synthèse mars 2019" in page


async def test_report_xlsx_applies_k_floor_and_direction(
    client: httpx.AsyncClient, staff: Staff, march_2019: dict[str, Any], fake_store: FakeStore
) -> None:
    headers = national(staff, ["ops_manager"])
    ar = await _create(client, headers, format="xlsx", kind="gouvernorat_summary", gouvernorats=[SILIANA, KEF])
    fr = await _create(client, headers, format="xlsx", lang="fr", filters={"category": "fire"})
    await _generate_all()
    book = load_workbook(io.BytesIO(fake_store.objects[f"analytics/reports/{ar['id']}.xlsx"]))
    cover = book.worksheets[0]
    assert cover.sheet_view.rightToLeft is True
    cover_text = " ".join(str(c.value) for row in cover.iter_rows() for c in row if c.value)
    assert "سليانة" in cover_text and "الكاف" in cover_text and "<5" in cover_text
    key_figures = book.worksheets[1]
    values = [row[1].value for row in key_figures.iter_rows(min_row=2)]
    # Siliana + Kef: 6 received (≥ 5, exact), every other cell under 5 is floored, never exact.
    assert values[0] == 6
    assert all(v in (0, "<5") or (isinstance(v, int) and v >= 5) for v in values)
    assert "<5" in values
    for sheet in book.worksheets[1:]:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                assert not (isinstance(cell.value, int) and 0 < cell.value < 5 and cell.column > 1), sheet.title
    book_fr = load_workbook(io.BytesIO(fake_store.objects[f"analytics/reports/{fr['id']}.xlsx"]))
    assert book_fr.worksheets[0].sheet_view.rightToLeft is False
    fr_cover = " ".join(str(c.value) for row in book_fr.worksheets[0].iter_rows() for c in row if c.value)
    assert "Incendie" in fr_cover and "k-anonymat" in fr_cover


async def test_report_library_scope_search_and_targets(
    client: httpx.AsyncClient, staff: Staff, march_2019: dict[str, Any]
) -> None:
    regional = staff(["local_committee"], SILIANA)
    token = f"zx-{uuid.uuid4().hex[:10]}"
    mine = await _create(client, regional, gouvernorats=[], title=f"Bilan feux Siliana {token}")
    assert mine["gouvernorats"] == [SILIANA]
    r = await client.post(
        "/api/v1/staff/reports-library",
        json={
            "kind": "period_summary",
            "title": "x",
            "period_from": "2019-03-01",
            "period_to": "2019-03-31",
            "gouvernorats": [KEF],
        },
        headers=regional,
    )
    assert r.status_code == 403 and r.json() == {"detail": "out_of_scope"}
    national_doc = await _create(client, national(staff), gouvernorats=[], title=f"National {token}")
    assert national_doc["gouvernorats"] == []
    # Visibility: the regional reader never sees national or other gouvernorats' documents.
    other = staff(["operator"], KEF)
    assert (await client.get(f"/api/v1/staff/reports-library/{mine['id']}", headers=other)).status_code == 404
    r = await client.get(f"/api/v1/staff/reports-library/{national_doc['id']}", headers=staff(["operator"], SILIANA))
    assert r.status_code == 404
    r = await client.get("/api/v1/staff/reports-library", params={"q": token}, headers=staff(["operator"], SILIANA))
    # The title match is always there; the events library index (when present) may add fuzzy matches, in scope only.
    items = r.json()["items"]
    assert mine["id"] in [i["id"] for i in items] and national_doc["id"] not in [i["id"] for i in items]
    assert all(i["gouvernorats"] == [SILIANA] for i in items)
    assert r.json()["search_backend"] in ("postgres", "knowledge")
    r = await client.get("/api/v1/staff/reports-library", params={"q": token}, headers=national(staff))
    assert {mine["id"], national_doc["id"]} <= {i["id"] for i in r.json()["items"]}
    r = await client.get(
        "/api/v1/staff/reports-library",
        params={"q": token, "kind": "period_summary", "from": "2019-03-15", "to": "2019-04-01", "gouvernorat": SILIANA},
        headers=national(staff),
    )
    items = r.json()["items"]
    assert mine["id"] in [i["id"] for i in items]
    assert all(i["kind"] == "period_summary" and SILIANA in i["gouvernorats"] for i in items)
    r = await client.get("/api/v1/staff/reports-library", params={"q": "50%_"}, headers=national(staff))
    assert r.status_code == 200
    # Not ready yet: no download.
    r = await client.get(f"/api/v1/staff/reports-library/{mine['id']}/download", headers=regional)
    assert r.status_code == 409 and r.json() == {"detail": "report_not_ready"}
    # Incident documents need an incident in scope; scenario documents a scenario.
    body = {"kind": "incident_summary", "title": "x", "period_from": "2019-03-01", "period_to": "2019-03-31"}
    r = await client.post("/api/v1/staff/reports-library", json=body, headers=regional)
    assert r.status_code == 422 and r.json() == {"detail": "incident_required"}
    r = await client.post(
        "/api/v1/staff/reports-library",
        json=body | {"filters": {"incident_id": str(march_2019["incident"])}},
        headers=staff(["local_committee"], KEF),
    )
    assert r.status_code == 403
    r = await client.post("/api/v1/staff/reports-library", json=body | {"kind": "scenario"}, headers=regional)
    assert r.status_code == 422 and r.json() == {"detail": "scenario_required"}
    r = await client.post(
        "/api/v1/staff/reports-library",
        json=body | {"period_to": "2019-02-01"},
        headers=regional,
    )
    assert r.status_code == 422 and r.json() == {"detail": "invalid_period"}
    r = await client.post("/api/v1/staff/reports-library", json=body, headers=staff(["operator"], SILIANA))
    assert r.status_code == 403  # operators read the library, managers and committees create


async def test_incident_after_action_and_scenario_documents(
    client: httpx.AsyncClient,
    staff: Staff,
    march_2019: dict[str, Any],
    tataouine: dict[str, Any],
    fake_store: FakeStore,
) -> None:
    headers = national(staff, ["central_committee"])
    incident = str(march_2019["incident"])
    after = await _create(client, headers, kind="after_action", gouvernorats=[], filters={"incident_id": incident})
    assert after["gouvernorats"] == [SILIANA]  # the incident's gouvernorat
    summary = await _create(
        client, headers, kind="incident_summary", format="xlsx", lang="fr", filters={"incident_id": incident}
    )
    run = await client.post("/api/v1/staff/analytics/scenarios", json=_scenario(), headers=headers)
    sim = await _create(
        client, headers, kind="scenario", lang="fr", gouvernorats=[], filters={"scenario_id": run.json()["id"]}
    )
    assert sim["gouvernorats"] == [TATAOUINE]
    await _generate_all()
    for doc in (after, summary, sim):
        r = await client.get(f"/api/v1/staff/reports-library/{doc['id']}", headers=headers)
        assert r.json()["status"] == "ready", r.json()
    assert fake_store.objects[f"analytics/reports/{after['id']}.pdf"].startswith(b"%PDF")
    assert fake_store.objects[f"analytics/reports/{sim['id']}.pdf"].startswith(b"%PDF")
    book = load_workbook(io.BytesIO(fake_store.objects[f"analytics/reports/{summary['id']}.xlsx"]))
    assert len(book.worksheets) == 4  # cover, card, reports, resources


async def test_report_fails_when_target_disappears(
    client: httpx.AsyncClient, staff: Staff, session: AsyncSession, march_2019: dict[str, Any]
) -> None:
    headers = national(staff, ["ops_manager"])
    incident = await seed_incident(session, GAFSA, "fire", local(date(2019, 3, 1)))
    await session.commit()
    doc = await _create(client, headers, kind="incident_summary", filters={"incident_id": str(incident)})
    await session.execute(text("DELETE FROM ops.incident WHERE id = :i"), {"i": incident})
    await session.commit()
    await _generate_all()
    r = await client.get(f"/api/v1/staff/reports-library/{doc['id']}", headers=headers)
    assert r.json()["status"] == "failed" and r.json()["error"] == "incident_not_found"


# ---------------------------------------------------------------- scheduled reports


async def test_schedule_crud_and_firing(
    client: httpx.AsyncClient,
    staff: Staff,
    make_token: TokenFactory,
    session: AsyncSession,
    fake_store: FakeStore,
) -> None:
    owner = str(uuid.uuid4())
    recipient = str(uuid.uuid4())
    headers = bearer(make_token(issuer=STAFF_ISS, roles=["local_committee"], gouvernorat_code=str(GAFSA), sub=owner))
    body = {
        "kind": "period_summary",
        "title": "Hebdo Gafsa",
        "format": "xlsx",
        "lang": "fr",
        "cron": {"frequency": "weekly", "time": "07:30", "weekday": 1},
        "recipients": [recipient.upper()],
    }
    for bad, code in (
        ({"cron": {"frequency": "weekly", "time": "07:30"}}, "invalid_cron"),
        ({"cron": {"frequency": "daily", "time": "25:00"}}, "invalid_cron"),
        ({"recipients": ["not-a-staff-id"]}, "invalid_recipients"),
        ({"filters": {"incident_id": str(uuid.uuid4())}}, "invalid_filters"),
    ):
        r = await client.post("/api/v1/staff/report-schedules", json=body | bad, headers=headers)
        assert r.status_code == 422 and r.json() == {"detail": code}, bad
    r = await client.post("/api/v1/staff/report-schedules", json=body | {"gouvernorats": [KEF]}, headers=headers)
    assert r.status_code == 403
    r = await client.post("/api/v1/staff/report-schedules", json=body, headers=headers)
    assert r.status_code == 201, r.text
    schedule = r.json()
    assert schedule["gouvernorats"] == [GAFSA] and schedule["recipients"] == [recipient]
    assert schedule["cron"] == {"frequency": "weekly", "time": "07:30", "weekday": 1, "day": None}
    next_run = datetime.fromisoformat(schedule["next_run_at"]).astimezone(TZ)
    assert next_run.isoweekday() == 1 and (next_run.hour, next_run.minute) == (7, 30)
    assert next_run > datetime.now(UTC)
    # Scope: another gouvernorat's manager does not see it; operators do not manage schedules.
    r = await client.get(f"/api/v1/staff/report-schedules/{schedule['id']}", headers=staff(["ops_manager"], KEF))
    assert r.status_code == 404
    assert (await client.get("/api/v1/staff/report-schedules", headers=staff(["operator"], GAFSA))).status_code == 403
    # Due now: the worker queues the document of the previous 7 days, generates it, notifies everyone.
    due = datetime(2026, 10, 19, 7, 30, tzinfo=TZ)  # a Monday
    await session.execute(
        text("UPDATE analytics.schedule SET next_run_at = :d WHERE id = :i"), {"d": due, "i": schedule["id"]}
    )
    await session.commit()
    await worker.cycle(get_sessionmaker())
    r = await client.get(f"/api/v1/staff/report-schedules/{schedule['id']}", headers=headers)
    fired = r.json()
    assert fired["last_run_at"] and fired["last_report_id"]
    assert datetime.fromisoformat(fired["next_run_at"]) > datetime.now(UTC)
    r = await client.get(f"/api/v1/staff/reports-library/{fired['last_report_id']}", headers=headers)
    report = r.json()
    assert report["status"] == "ready" and report["schedule_id"] == schedule["id"]
    assert (report["period_from"], report["period_to"]) == ("2026-10-12", "2026-10-18")
    assert report["gouvernorats"] == [GAFSA] and report["format"] == "xlsx"
    assert f"analytics/reports/{report['id']}.xlsx" in fake_store.objects
    # Both the recipient and the owner have an in-app notification; marking it read is per person.
    reader = bearer(make_token(issuer=STAFF_ISS, roles=["operator"], gouvernorat_code=str(GAFSA), sub=recipient))
    r = await client.get("/api/v1/staff/analytics/notifications", headers=reader)
    items = r.json()["items"]
    assert r.json()["unread"] == 1 and items[0]["schedule_id"] == schedule["id"] and items[0]["kind"] == "report_ready"
    r = await client.post(f"/api/v1/staff/analytics/notifications/{items[0]['id']}/read", headers=reader)
    assert r.status_code == 204
    assert (await client.get("/api/v1/staff/analytics/notifications", headers=reader)).json()["unread"] == 0
    assert (await client.get("/api/v1/staff/analytics/notifications", headers=headers)).json()["unread"] == 1
    r = await client.post(f"/api/v1/staff/analytics/notifications/{items[0]['id']}/read", headers=headers)
    assert r.status_code == 404  # someone else's notification
    run_audit = await session.scalar(
        text("SELECT count(*) FROM audit.log WHERE action = 'analytics.schedule.run' AND target = :t"),
        {"t": f"analytics_schedule:{schedule['id']}"},
    )
    assert run_audit == 1
    # A second cycle does not fire again before the next run.
    await worker.cycle(get_sessionmaker())
    count = await session.scalar(
        text("SELECT count(*) FROM analytics.report WHERE schedule_id = :i"), {"i": schedule["id"]}
    )
    assert count == 1
    # Update (monthly, disabled) and delete; the produced document stays.
    r = await client.patch(
        f"/api/v1/staff/report-schedules/{schedule['id']}",
        json={"cron": {"frequency": "monthly", "time": "06:00", "day": 1}, "enabled": False},
        headers=headers,
    )
    assert r.status_code == 200 and r.json()["enabled"] is False and r.json()["cron"]["day"] == 1
    r = await client.delete(f"/api/v1/staff/report-schedules/{schedule['id']}", headers=headers)
    assert r.status_code == 204
    r = await client.get(f"/api/v1/staff/reports-library/{report['id']}", headers=headers)
    assert r.status_code == 200 and r.json()["schedule_id"] is None
    actions = set(
        await session.scalars(
            text("SELECT action FROM audit.log WHERE target = :t"), {"t": f"analytics_schedule:{schedule['id']}"}
        )
    )
    assert {"analytics.schedule.create", "analytics.schedule.update", "analytics.schedule.delete"} <= actions


# ---------------------------------------------------------------- phase 9: English, participation, national ranking


async def test_document_in_english(
    client: httpx.AsyncClient, staff: Staff, march_2019: dict[str, Any], fake_store: FakeStore
) -> None:
    created = await _create(client, national(staff, ["ops_manager"]), lang="en", title="March 2019 summary")
    await _generate_all()
    r = await client.get(f"/api/v1/staff/reports-library/{created['id']}", headers=national(staff, ["ops_manager"]))
    assert r.json()["status"] == "ready", r.json()
    assert fake_store.objects[f"analytics/reports/{created['id']}.pdf"].startswith(b"%PDF")
    from app.modules.analytics import documents
    from app.modules.analytics.models import Report

    async with get_sessionmaker()() as s:
        row = await s.get(Report, uuid.UUID(created["id"]))
        assert row is not None
        page = documents.render_html(await documents.build(s, row))
    assert 'dir="ltr"' in page and 'lang="en"' in page
    assert "Reports received" in page and "Signalements reçus" not in page


async def test_participating_breakdown(
    client: httpx.AsyncClient, staff: Staff, march_2019: dict[str, Any], session: AsyncSession
) -> None:
    """An incident led by Kef with Siliana participating counts once for Kef (lead) and once, as participating, for
    Siliana; the Siliana-led incident of the fixture is not participating."""
    shared = await seed_incident(session, KEF, "fire", local(date(2019, 3, 20), 8))
    await session.execute(
        text(
            "INSERT INTO ops.incident_gouvernorat (incident_id, gouvernorat_code, role, added_by, reason) "
            "VALUES (:i, :g, 'participating', 'test', 'renfort')"
        ),
        {"i": shared, "g": SILIANA},
    )
    await session.commit()
    params = march_2019["params"]
    siliana = (
        await client.get("/api/v1/staff/analytics/overview", params=params, headers=staff(["operator"], SILIANA))
    ).json()
    assert {(i["participating"], i["count"]) for i in siliana["incidents"]} == {(False, 1), (True, 1)}
    kef = (await client.get("/api/v1/staff/analytics/overview", params=params, headers=staff(["operator"], KEF))).json()
    assert {(i["participating"], i["count"]) for i in kef["incidents"]} == {(False, 1)}
    # The national room reads the same split per gouvernorat.
    national_kef = (
        await client.get(
            "/api/v1/staff/analytics/overview", params=params | {"gouvernorat": KEF}, headers=national(staff)
        )
    ).json()
    assert national_kef["incidents"] == kef["incidents"]


async def test_leaderboard_national_refused(client: httpx.AsyncClient, staff: Staff) -> None:
    url = "/api/v1/staff/rewards/leaderboard"
    r = await client.get(url, params={"scope": "national"}, headers=staff(["ops_manager"], SILIANA))
    assert (r.status_code, r.json()["detail"]) == (403, "out_of_scope")
    r = await client.get(url, params={"gouvernorat": KEF}, headers=staff(["ops_manager"], SILIANA))
    assert (r.status_code, r.json()["detail"]) == (403, "out_of_scope")
    assert (await client.get(url, headers=staff(["ops_manager"], SILIANA))).status_code == 200
    r = await client.get(url, params={"scope": "national"}, headers=national(staff, ["ops_manager"]))
    assert r.status_code == 200 and r.json()["scope"] == "national"
