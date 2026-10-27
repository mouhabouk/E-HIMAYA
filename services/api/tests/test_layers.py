"""Map data layers (phase B1, docs/phase-b/design.md §2): file reading, preview, apply, undo, export, public reads."""

import csv
import io
import json
import uuid
from collections.abc import Callable
from typing import Any

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.layers import domain

Staff = Callable[..., dict[str, str]]
URL = "/api/v1/staff/layers"

SFAX = (34.7406, 10.7603)
SOUSSE = (35.8256, 10.6360)
TUNIS = (36.8065, 10.1815)
KAIROUAN = (35.6781, 10.0963)


# ---------------------------------------------------------------- reading files (no database)


def test_csv_semicolon_arabic_headings_decimal_comma_and_phones() -> None:
    data = (
        "المعرف;الاسم;النوع;خط العرض;خط الطول;الهاتف;المسؤول;عدد الأسرة\n"
        "H1;مستشفى الهادي شاكر;مستشفى;34,7406;10,7603;74 241 511 / 74 241 512;د. فلان;420\n"
    ).encode()
    parsed = domain.parse("hopitaux.csv", data)
    assert parsed.format == "csv" and not parsed.rejected
    row = parsed.rows[0]
    assert (row.ref, row.name, row.kind, row.lat, row.lon) == ("H1", "مستشفى الهادي شاكر", "مستشفى", 34.7406, 10.7603)
    assert row.phone == "+21674241511"
    assert row.props == {"phone2": "+21674241512", "manager": "د. فلان", "عدد الأسرة": "420"}
    assert row.hidden is None


def test_csv_comma_french_headings_bom_and_rejections() -> None:
    data = (
        "﻿identifiant,Nom,Latitude,Longitude,Téléphone,E-mail,Site web\n"
        "A,Clinique A,36.8065,10.1815,198,contact@a.tn,www.a.tn\n"
        ",Sans position,,,,,\n"
        "C,Inversée,10.18,36.80,,,\n"
        "D,,36.8,10.1,,,\n"
        "A,Doublon,36.8,10.1,,,\n"
        "E,Hors pays,48.85,2.35,12,pas-un-mail,javascript:alert(1)\n"
        ",,,,,,\n"
    ).encode()
    parsed = domain.parse("cliniques.csv", data)
    assert [(r.ref, r.phone, r.props) for r in parsed.rows] == [
        ("A", "198", {"email": "contact@a.tn", "website": "https://www.a.tn"})
    ]
    assert [(r.line, r.reason) for r in parsed.rejected] == [
        (3, "coordinates_missing"),
        (4, "coordinates_swapped"),
        (5, "name_missing"),
        (6, "duplicate_id"),
        (7, "outside_tunisia"),
    ]


def test_row_warnings_keep_the_place() -> None:
    data = b"name,lat,lon,phone,email,website\nP,36.8,10.18,12,bad,javascript:alert(1)\n"
    parsed = domain.parse("p.csv", data)
    assert len(parsed.rows) == 1 and parsed.rows[0].phone is None and parsed.rows[0].props == {}
    assert sorted(w.reason for w in parsed.warnings) == ["email_invalid", "phone_invalid", "website_invalid"]


def test_windows_arabic_code_page() -> None:
    data = "الاسم,خط العرض,خط الطول\nصيدلية,36.8,10.18\n".encode("cp1256")
    assert domain.parse("p.csv", data).rows[0].name == "صيدلية"


def test_xlsx_numbers_and_english_headings() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["id", "name", "lat", "lon", "phone", "hidden", "beds"])
    sheet.append([101, "Hospital", 35.8256, 10.636, 73241411, "yes", 300])
    sheet.append([102, '=HYPERLINK("x")', 35.8, 10.6, None, None, None])
    buffer = io.BytesIO()
    book.save(buffer)
    parsed = domain.parse("h.xlsx", buffer.getvalue())
    assert parsed.format == "xlsx"
    (first,) = parsed.rows
    assert (first.ref, first.phone, first.hidden, first.props) == ("101", "+21673241411", True, {"beds": "300"})
    # Formulas are never evaluated: without a value saved by Excel the cell is empty.
    assert [(r.line, r.reason) for r in parsed.rejected] == [(3, "name_missing")]


def test_json_and_geojson() -> None:
    items = [{"nom": "Abri 1", "lat": 35.67, "lon": 10.09, "capacité": 120}]
    assert domain.parse("abris.json", json.dumps(items).encode()).rows[0].props == {"capacité": "120"}
    doc = {
        "type": "FeatureCollection",
        "features": [
            {
                "type": "Feature",
                "id": "node/1",
                "geometry": {"type": "Point", "coordinates": [10.18, 36.8]},
                "properties": {"name": "Pharmacie", "opening_hours": "24/7"},
            },
            {
                "type": "Feature",
                "geometry": {"type": "LineString", "coordinates": [[10, 36], [10.1, 36.1]]},
                "properties": {"name": "Route"},
            },
        ],
    }
    parsed = domain.parse("osm.geojson", json.dumps(doc).encode())
    assert parsed.format == "geojson"
    point, line = parsed.rows
    assert (point.ref, point.lat, point.lon, point.props) == ("node/1", 36.8, 10.18, {"hours": "24/7"})
    # Since phase L1a a line is kept whole; its point is one on the line.
    assert line.name == "Route" and line.geometry == "line" and line.shape is not None
    assert parsed.rejected == []


@pytest.mark.parametrize(
    ("name", "data", "code"),
    [
        ("a.csv", b"", "empty_file"),
        ("a.csv", b"foo,bar\n1,2\n", "required_columns_missing"),
        ("a.json", b"{not json", "json_invalid"),
        ("a.json", b'{"a": 1}', "json_not_array"),
        ("a.geojson", b'{"type": "Feature"}', "geojson_not_collection"),
        ("a.xlsx", b"PK\x03\x04garbage", "xlsx_unreadable"),
        ("a.pdf", b"%PDF-1.7", "format_unknown"),
        ("a.csv", b"x" * (domain.MAX_FILE_BYTES + 1), "file_too_large"),
    ],
)
def test_refused_files(name: str, data: bytes, code: str) -> None:
    with pytest.raises(domain.FileError) as e:
        domain.parse(name, data)
    assert e.value.code == code


def test_too_many_rows() -> None:
    data = "name,lat,lon\n" + "p,36.8,10.18\n" * (domain.MAX_ROWS + 1)
    with pytest.raises(domain.FileError, match="too_many_rows"):
        domain.parse("big.csv", data.encode())


# ---------------------------------------------------------------- API helpers


def _key() -> str:
    return f"t_{uuid.uuid4().hex[:10]}"


def _admin(staff: Staff) -> dict[str, str]:
    return staff(["sysadmin"], gouvernorat=None)


async def _create(client: httpx.AsyncClient, staff: Staff, audience: str = "public", **extra: Any) -> dict[str, Any]:
    body = {
        "key": _key(),
        "name": {"ar": "المستشفيات", "fr": "Hôpitaux", "en": "Hospitals"},
        "icon": "hospital",
        "colour": "#c62828",
        "audience": audience,
        "source": "OpenStreetMap",
        "licence": "ODbL",
        "reason": "first load of the layer",
        # Since phase L1a a new layer starts as a draft; these tests use it published.
        "status": "published",
        **extra,
    }
    r = await client.post(URL, json=body, headers=_admin(staff))
    assert r.status_code == 201, r.text
    return r.json()


def _csv(rows: list[list[Any]], header: list[str] | None = None) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header or ["id", "name", "type", "lat", "lon", "phone", "manager", "beds"])
    writer.writerows(rows)
    return buffer.getvalue().encode()


HOSPITALS = [
    ["S1", "Hôpital Habib Bourguiba", "hospital", *SFAX, "74241511", "Dr A", "420"],
    ["S2", "Hôpital Sahloul", "hospital", *SOUSSE, "73369411", "Dr B", "600"],
    ["T1", "Hôpital Charles Nicolle", "emergency", *TUNIS, "71578000", "Dr C", ""],
]


async def _preview(
    client: httpx.AsyncClient, headers: dict[str, str], key: str, data: bytes, mode: str = "merge", name: str = "f.csv"
) -> httpx.Response:
    return await client.post(
        f"{URL}/{key}/imports",
        params={"file_name": name, "mode": mode},
        content=data,
        headers={**headers, "Content-Type": "application/octet-stream"},
    )


async def _import(
    client: httpx.AsyncClient, headers: dict[str, str], key: str, data: bytes, mode: str = "merge"
) -> dict[str, Any]:
    r = await _preview(client, headers, key, data, mode)
    assert r.status_code == 201, r.text
    applied = await client.post(
        f"{URL}/{key}/imports/{r.json()['id']}/apply", json={"reason": "official list"}, headers=headers
    )
    assert applied.status_code == 200, applied.text
    return r.json()


# ---------------------------------------------------------------- API


async def test_who_may_create_and_import(client: httpx.AsyncClient, staff: Staff, citizen: dict[str, str]) -> None:
    body = {
        "key": _key(),
        "name": {"ar": "أ", "fr": "A", "en": "A"},
        "icon": "pin",
        "colour": "#2E7D32",
        "reason": "test layer",
    }
    assert (await client.post(URL, json=body)).status_code == 401
    assert (await client.post(URL, json=body, headers=citizen)).status_code == 403
    r = await client.post(URL, json=body, headers=staff(["sysadmin"], gouvernorat=61))
    assert (r.status_code, r.json()["detail"]) == (403, "national_scope_required")
    r = await client.post(URL, json=body, headers=staff(["content_editor"], gouvernorat=None))
    assert (r.status_code, r.json()["detail"]) == (403, "insufficient_role")
    r = await client.post(URL, json={**body, "colour": "#FFB300"}, headers=_admin(staff))
    assert r.status_code == 422  # amber is reserved for simulations
    r = await client.post(URL, json={**body, "key": "civil_protection"}, headers=_admin(staff))
    assert (r.status_code, r.json()["detail"]) == (409, "key_taken")

    layer = await _create(client, staff)
    assert layer["staff_only"] == ["manager"] and layer["colour"] == "#C62828"
    editor = staff(["content_editor"], gouvernorat=None)
    r = await _preview(client, editor, layer["key"], _csv(HOSPITALS))
    assert (r.status_code, r.json()["detail"]) == (403, "insufficient_role")
    r = await client.patch(
        f"{URL}/{layer['key']}",
        json={"version": layer["version"], "editors_allowed": True, "reason": "editors keep it up to date"},
        headers=_admin(staff),
    )
    assert r.status_code == 200 and r.json()["editors_allowed"] is True
    listed = {i["key"]: i for i in (await client.get(URL, headers=editor)).json()["items"]}
    assert listed[layer["key"]]["can_import"] is True and listed[layer["key"]]["can_manage"] is False
    assert (await _preview(client, editor, layer["key"], _csv(HOSPITALS))).status_code == 201
    # A regional operator reads the layers but imports nothing.
    regional = staff(["operator"], gouvernorat=61)
    assert (await client.get(URL, headers=regional)).status_code == 200
    assert (await _preview(client, regional, layer["key"], _csv(HOSPITALS))).status_code == 403


async def test_preview_apply_public_and_staff_only_columns(
    client: httpx.AsyncClient, staff: Staff, session: AsyncSession
) -> None:
    layer = await _create(client, staff)
    key, admin = layer["key"], _admin(staff)
    rows = [*HOSPITALS, ["X1", "Hors frontière", "", 41.0, 9.0, "", "", ""]]
    r = await _preview(client, admin, key, _csv(rows))
    assert r.status_code == 201, r.text
    preview = r.json()
    assert preview["status"] == "preview" and preview["can_apply"] is True
    assert {k: preview["counts"][k] for k in ("new", "updated", "unchanged", "removed", "rejected")} == {
        "new": 3,
        "updated": 0,
        "unchanged": 0,
        "removed": 0,
        "rejected": 1,
    }
    assert preview["rejected"] == [{"line": 5, "reason": "outside_tunisia", "name": "Hors frontière"}]
    assert {p["ref"] for p in preview["points"]} == {"S1", "S2", "T1"}
    # Nothing on the map before it is applied.
    near = f"{SFAX[0]},{SFAX[1]}"
    r = await client.get(f"/api/v1/layers/{key}/features", params={"near": near})
    assert r.json()["items"] == []

    r = await client.post(f"{URL}/{key}/imports/{preview['id']}/apply", json={"reason": "x"}, headers=admin)
    assert r.status_code == 422
    r = await client.post(
        f"{URL}/{key}/imports/{preview['id']}/apply", json={"reason": "OSM first load"}, headers=admin
    )
    assert r.status_code == 200 and r.json()["status"] == "applied" and r.json()["can_undo"] is True

    # Public: nearest first, with the distance; the manager's name never leaves the API.
    r = await client.get(f"/api/v1/layers/{key}/features", params={"near": near, "limit": 3, "radius_km": 300})
    items = r.json()["items"]
    assert [i["name"] for i in items] == ["Hôpital Habib Bourguiba", "Hôpital Sahloul", "Hôpital Charles Nicolle"]
    assert items[0]["distance_m"] < 1 and items[1]["distance_m"] > 100_000
    assert items[0]["phone"] == "+21674241511" and items[0]["props"] == {"beds": "420"}
    assert items[0]["gouvernorat_code"] == 61
    assert all("manager" not in i["props"] for i in items)
    r = await client.get(f"/api/v1/layers/{key}/features", params={"near": near, "radius_km": 50})
    assert len(r.json()["items"]) == 1
    # Staff see every column.
    r = await client.get(f"{URL}/{key}/features", headers=staff())
    assert {f["ref"]: f["props"].get("manager") for f in r.json()["items"]} == {
        "S1": "Dr A",
        "S2": "Dr B",
        "T1": "Dr C",
    }

    entries = (
        await session.execute(
            text("SELECT action, reason FROM audit.log WHERE target = :t ORDER BY seq"), {"t": f"layer:{key}"}
        )
    ).all()
    assert [e.action for e in entries] == ["layer.create", "layer.import_apply"]
    assert entries[1].reason == "OSM first load"


async def test_merge_replace_and_undo(client: httpx.AsyncClient, staff: Staff) -> None:
    layer = await _create(client, staff)
    key, admin = layer["key"], _admin(staff)
    await _import(client, admin, key, _csv(HOSPITALS))

    # Merge: one changed, one new, the others kept.
    changed = [["S1", "Hôpital Habib Bourguiba", "hospital", *SFAX, "74241511", "Dr A2", "450"]]
    added = [["K1", "Hôpital Ibn El Jazzar", "hospital", *KAIROUAN, "77226200", "", ""]]
    r = await _preview(client, admin, key, _csv(changed + added))
    counts = r.json()["counts"]
    assert (counts["new"], counts["updated"], counts["unchanged"], counts["removed"]) == (1, 1, 0, 0)
    await client.post(f"{URL}/{key}/imports/{r.json()['id']}/apply", json={"reason": "update"}, headers=admin)
    names = await _staff_refs(client, staff, key)
    assert names == {"S1", "S2", "T1", "K1"}

    # Replace: the file becomes the whole layer.
    r = await _preview(client, admin, key, _csv(HOSPITALS[:2]), mode="replace")
    counts = r.json()["counts"]
    assert (counts["new"], counts["updated"], counts["unchanged"], counts["removed"]) == (0, 1, 1, 2)
    applied = await client.post(
        f"{URL}/{key}/imports/{r.json()['id']}/apply", json={"reason": "official list"}, headers=admin
    )
    assert await _staff_refs(client, staff, key) == {"S1", "S2"}

    # Undo the last import: the four places come back as they were.
    undo = await client.post(
        f"{URL}/{key}/imports/{applied.json()['id']}/undo", json={"reason": "wrong file"}, headers=admin
    )
    assert undo.status_code == 200 and undo.json()["status"] == "undone"
    assert await _staff_refs(client, staff, key) == {"S1", "S2", "T1", "K1"}
    features = (await client.get(f"{URL}/{key}/features", headers=staff())).json()["items"]
    s1 = next(f for f in features if f["ref"] == "S1")
    assert s1["props"]["manager"] == "Dr A2" and s1["props"]["beds"] == "450"
    # Only the last applied import can be undone, once.
    again = await client.post(
        f"{URL}/{key}/imports/{applied.json()['id']}/undo", json={"reason": "wrong file"}, headers=admin
    )
    assert (again.status_code, again.json()["detail"]) == (409, "not_last_import")


async def _staff_refs(client: httpx.AsyncClient, staff: Staff, key: str) -> set[str]:
    r = await client.get(f"{URL}/{key}/features", headers=staff())
    return {f["ref"] for f in r.json()["items"]}


async def test_stale_preview_and_change_after_import(client: httpx.AsyncClient, staff: Staff) -> None:
    layer = await _create(client, staff)
    key, admin = layer["key"], _admin(staff)
    first = await _preview(client, admin, key, _csv(HOSPITALS))
    applied = await _import(client, admin, key, _csv(HOSPITALS[:1]))
    # The layer changed after the first preview: preview again.
    r = await client.post(
        f"{URL}/{key}/imports/{first.json()['id']}/apply", json={"reason": "late file"}, headers=admin
    )
    assert (r.status_code, r.json()["detail"]) == (409, "layer_changed")
    # A correction by hand after an import blocks its undo.
    feature = (await client.get(f"{URL}/{key}/features", headers=staff())).json()["items"][0]
    r = await client.patch(
        f"{URL}/{key}/features/{feature['id']}",
        json={"version": feature["version"], "hidden": True, "reason": "closed for works"},
        headers=admin,
    )
    assert r.status_code == 200 and r.json()["hidden"] is True
    r = await client.post(f"{URL}/{key}/imports/{applied['id']}/undo", json={"reason": "undo it"}, headers=admin)
    assert (r.status_code, r.json()["detail"]) == (409, "layer_changed")
    # A hidden place leaves the public map, stays on the staff one.
    r = await client.get(f"/api/v1/layers/{key}/features", params={"near": f"{SFAX[0]},{SFAX[1]}"})
    assert r.json()["items"] == []
    # A file without a "hidden" column does not bring it back.
    r = await _preview(client, admin, key, _csv(HOSPITALS[:1]))
    assert r.json()["counts"]["unchanged"] == 1


async def test_correction_by_hand(client: httpx.AsyncClient, staff: Staff) -> None:
    layer = await _create(client, staff)
    key, admin = layer["key"], _admin(staff)
    body = {"name": "Poste avancé", "lat": SOUSSE[0], "lon": SOUSSE[1], "phone": "73 000 000", "reason": "new site"}
    r = await client.post(f"{URL}/{key}/features", json=body, headers=admin)
    assert r.status_code == 201, r.text
    created = r.json()
    assert created["phone"] == "+21673000000" and created["gouvernorat_code"] == 51
    assert (await client.post(f"{URL}/{key}/features", json=body, headers=admin)).json()["detail"] == "duplicate_id"
    r = await client.patch(
        f"{URL}/{key}/features/{created['id']}",
        json={"version": created["version"], "lat": KAIROUAN[0], "lon": KAIROUAN[1], "reason": "moved"},
        headers=admin,
    )
    assert r.status_code == 200 and r.json()["gouvernorat_code"] == 41
    r = await client.patch(
        f"{URL}/{key}/features/{created['id']}",
        json={"version": created["version"], "name": "Old", "reason": "stale edit"},
        headers=admin,
    )
    assert (r.status_code, r.json()["detail"]) == (409, "version_conflict")
    r = await client.post(f"{URL}/{key}/features", json={**body, "phone": "12"}, headers=admin)
    assert r.status_code == 422


async def test_export_then_import_gives_the_same_layer(client: httpx.AsyncClient, staff: Staff) -> None:
    layer = await _create(client, staff)
    key, admin = layer["key"], _admin(staff)
    rows = [*HOSPITALS, ["F1", "=cmd|' /C calc'!A0", "", *KAIROUAN, "", "", ""]]
    await _import(client, admin, key, _csv(rows))
    # Since phase L1a the whole file (every column) is for the administrator and the layer's editors only.
    assert (await client.get(f"{URL}/{key}/export", headers=staff())).status_code == 403
    for fmt, lang in (("csv", "ar"), ("xlsx", "fr"), ("geojson", "en")):
        r = await client.get(f"{URL}/{key}/export", params={"format": fmt, "lang": lang}, headers=admin)
        assert r.status_code == 200, r.text
        assert f"{key}-v1.{fmt}" in r.headers["content-disposition"]
        again = await _preview(client, admin, key, r.content, name=f"export.{fmt}")
        assert again.status_code == 201, again.text
        counts = again.json()["counts"]
        if fmt == "geojson":
            assert (counts["unchanged"], counts["new"], counts["updated"]) == (4, 0, 0)
        else:
            # The neutralised cell comes back with its leading quote: the one place reported as updated.
            assert (counts["unchanged"], counts["new"], counts["updated"]) == (3, 0, 1)
    csv_text = (await client.get(f"{URL}/{key}/export", headers=admin)).content.decode("utf-8-sig")
    assert "'=cmd" in csv_text and "+21674241511" in csv_text and "'+216" not in csv_text
    xlsx = await client.get(f"{URL}/{key}/export", params={"format": "xlsx"}, headers=admin)
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet is not None and sheet["A1"].value == "identifiant"


async def test_template(client: httpx.AsyncClient, staff: Staff) -> None:
    r = await client.get(f"{URL}/template", params={"lang": "ar"}, headers=staff())
    assert r.status_code == 200
    header = r.content.decode("utf-8-sig").splitlines()[0].split(",")
    assert header[:5] == ["المعرف", "الاسم", "النوع", "خط العرض", "خط الطول"]
    assert all(domain.column_key(h) for h in header)
    r = await client.get(f"{URL}/template", params={"format": "xlsx", "lang": "fr"}, headers=staff())
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")


async def test_public_list_package_and_staff_layers(client: httpx.AsyncClient, staff: Staff) -> None:
    public = await _create(client, staff)
    hidden_layer = await _create(client, staff, audience="staff")
    admin = _admin(staff)
    await _import(client, admin, public["key"], _csv(HOSPITALS))
    await _import(client, admin, hidden_layer["key"], _csv(HOSPITALS))

    r = await client.get("/api/v1/layers", params={"lang": "fr"})
    assert r.status_code == 200
    items = {i["key"]: i for i in r.json()["items"]}
    assert r.json()["items"][0]["key"] == "civil_protection" and items["civil_protection"]["origin"] == "units"
    assert items[public["key"]]["name"] == "Hôpitaux" and items[public["key"]]["count"] == 3
    assert hidden_layer["key"] not in items
    etag = r.headers["etag"]
    assert (
        await client.get("/api/v1/layers", params={"lang": "fr"}, headers={"If-None-Match": etag})
    ).status_code == 304
    assert (await client.get(f"/api/v1/layers/{hidden_layer['key']}/package")).status_code == 404

    r = await client.get(f"/api/v1/layers/{public['key']}/package", params={"gouvernorat": 61})
    assert [i["name"] for i in r.json()["items"]] == ["Hôpital Habib Bourguiba"]
    etag = r.headers["etag"]
    same = await client.get(
        f"/api/v1/layers/{public['key']}/package", params={"gouvernorat": 61}, headers={"If-None-Match": etag}
    )
    assert same.status_code == 304
    # A change of the layer gives a new package.
    await _import(client, admin, public["key"], _csv([["S1", "Hôpital HB", "hospital", *SFAX, "", "", ""]]))
    changed = await client.get(
        f"/api/v1/layers/{public['key']}/package", params={"gouvernorat": 61}, headers={"If-None-Match": etag}
    )
    assert changed.status_code == 200 and changed.json()["items"][0]["name"] == "Hôpital HB"

    # Turning a layer off takes it out of the app.
    current = next(i for i in (await client.get(URL, headers=admin)).json()["items"] if i["key"] == public["key"])
    r = await client.patch(
        f"{URL}/{public['key']}",
        json={"version": current["version"], "enabled": False, "reason": "data under review"},
        headers=admin,
    )
    assert r.status_code == 200
    r = await client.get(f"/api/v1/layers/{public['key']}/features", params={"near": "34.7,10.7"})
    assert r.status_code == 404


async def test_public_query_checks(client: httpx.AsyncClient, staff: Staff) -> None:
    layer = await _create(client, staff)
    base = f"/api/v1/layers/{layer['key']}/features"
    assert (await client.get(base)).json()["detail"] == "near_or_bbox_required"
    assert (await client.get(base, params={"near": "abc"})).json()["detail"] == "invalid_near"
    assert (await client.get(base, params={"bbox": "11,37,10,36"})).json()["detail"] == "invalid_bbox"
    await _import(client, _admin(staff), layer["key"], _csv(HOSPITALS))
    r = await client.get(base, params={"bbox": "10.0,36.5,10.5,37.0"})
    assert [i["name"] for i in r.json()["items"]] == ["Hôpital Charles Nicolle"]


async def test_civil_protection_centres_come_from_the_units(client: httpx.AsyncClient, session: AsyncSession) -> None:
    async with session.begin():
        await session.execute(
            text(
                "INSERT INTO ops.unit (name_ar, name_fr, kind, gouvernorat_code, location, phone) VALUES "
                "('وحدة اختبار', 'Unité test', 'unit', 61, 'SRID=4326;POINT(10.7604 34.7407)', '+21674000000')"
            )
        )
    r = await client.get(
        "/api/v1/layers/civil_protection/features", params={"near": "34.7406,10.7603", "limit": 1, "lang": "ar"}
    )
    assert r.status_code == 200, r.text
    item = r.json()["items"][0]
    assert item["name"] == "وحدة اختبار" and item["distance_m"] < 50 and item["phone"] == "+21674000000"


async def test_file_errors_through_the_api(client: httpx.AsyncClient, staff: Staff) -> None:
    layer = await _create(client, staff)
    admin = _admin(staff)
    r = await _preview(client, admin, layer["key"], b"foo,bar\n1,2\n")
    assert (r.status_code, r.json()["detail"]) == (422, "required_columns_missing")
    r = await _preview(client, admin, layer["key"], b"x" * (domain.MAX_FILE_BYTES + 1))
    assert (r.status_code, r.json()["detail"]) == (413, "file_too_large")
    r = await _preview(client, admin, "missing_layer", _csv(HOSPITALS))
    assert r.status_code == 404
